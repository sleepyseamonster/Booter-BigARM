"""Apply the October 16 value audit with a preserving, revision-checked merge.

The JSON audit owns the reviewed assignments. This is catalog-only authoring:
it creates no Unity objects, recipes, prices, contents, energy or save records.
"""
import argparse
import json
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from category_edits import parts_for, patch_cells, verify_move
from tier_edits import INGREDIENTS, tier_number
from workbook_edits import atomic_write, digest, master_items, parse_sheets

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
AUDIT = ROOT / 'Docs/Design/Items/SALVAGE_VALUE_TIERS_2026-10-16.json'
OUT = ROOT / 'Logs/SalvageTiers20261016'


def prepare():
    snapshot = SOURCE.read_bytes()
    sheets = parse_sheets(snapshot)
    items = [(n, dict(zip(sheets['Items'][4], r))) for n, r in master_items(sheets)]
    salvage = [(n, r) for n, r in items if r['Category'] == 'Salvage']
    assessment = [line.split('|', 2) for line in (OUT / 'assessment.txt').read_text().splitlines()]
    assert len(salvage) == len(assessment) == 200
    assert [int(a[0]) for a in assessment] == list(range(1, 201))
    assert Counter(a[1] for a in assessment) == {str(t): 40 for t in range(5)}
    assert len({r['Item ID'] for _, r in items}) == len(items)
    selected = {}
    reviewed = []
    for (n, row), (_, tier, reason) in zip(salvage, assessment):
        value = 'T' + tier
        selected[row['Item ID']] = value
        reviewed.append(dict(id=row['Item ID'], name=row['Item name'], row=n,
                             previousTier=row['Tier'], tier=value, valueReason=reason))
    changes = {}

    def set_cell(sheet, n, field, value):
        header = sheets[sheet][4]
        ci = header.index(field)
        row = list(sheets[sheet][n-1])
        if row[ci] != value:
            changes.setdefault(sheet, {})[f'{openpyxl.utils.get_column_letter(ci+1)}{n}'] = value
            row[ci] = value
            sheets[sheet][n-1] = tuple(row)

    for item in reviewed:
        n, tier = item['row'], item['tier']
        set_cell('Items', n, 'Tier', tier)
        set_cell('Items', n, 'Tier basis', f"Salvage value audit October 16, 2026: {tier}. {item['valueReason']} Value tier is relative to this 200-item salvage roster; condition and crafting qualification remain separate.")
        set_cell('Items', n, 'Tier sort order', tier_number(tier))
        row = dict(zip(sheets['Items'][4], sheets['Items'][n-1]))
        fields = ['Item ID','Item name','Category','Subtype','Item type','Processing stage','Origin family','Primary use','Material family','Physical form','Design status','Search aliases','Tier','Finished-item input tier','Crafting role']
        set_cell('Items', n, 'Search terms', '; '.join(str(row.get(k) or '') for k in fields).lower())
        coverage = ('Route listed; review status/completeness' if (row.get('Listed output routes') or 0) > 0
                    else 'No production route listed')
        set_cell('Items', n, 'Crafting coverage', coverage)

    linked = {'Tier review': [('Item ID','Current tier'), ('Item ID','Tier sort order')],
              'Craft audit': [('Item ID','Tier')], 'Material routes': [('Item ID','Material tier')],
              'Crafting browser': [('Input item ID','Input tier'), ('Output item ID','Output tier')]}
    for sheet, mappings in linked.items():
        h = sheets[sheet][4]
        for n, row in enumerate(sheets[sheet][5:], 6):
            for key, field in mappings:
                key_value = row[h.index(key)]
                if key_value in selected:
                    tier = selected[key_value]
                    set_cell(sheet, n, field, tier_number(tier) if field == 'Tier sort order' else tier)
    reasons = {i['id']:i['valueReason'] for i in reviewed}
    h = sheets['Tier review'][4]
    for n, row in enumerate(sheets['Tier review'][5:], 6):
        if row[h.index('Item ID')] in selected:
            set_cell('Tier review', n, 'Decision basis', 'October 16 salvage value assignment. ' + reasons[row[h.index('Item ID')]] + ' Crafting qualification remains separate.')
    recipes = set()
    rh = sheets['Material routes'][4]
    for n, row in enumerate(sheets['Material routes'][5:], 6):
        if row[rh.index('Item ID')] in selected:
            ingredient = row[rh.index('Relationship')] in INGREDIENTS
            numeric = tier_number(row[rh.index('Material tier')]) if ingredient else -1
            set_cell('Material routes', n, 'Ingredient tier number', numeric)
            set_cell('Material routes', n, 'Unresolved ingredient tier', int(ingredient and numeric == -1))
            if ingredient:
                recipes.add(row[rh.index('Recipe ID')])
    listed = {}
    ch = sheets['Crafting recipes'][4]
    for n, row in enumerate(sheets['Crafting recipes'][5:], 6):
        key = row[ch.index('Recipe ID')]
        if key in recipes:
            nums = [tier_number(r[rh.index('Material tier')]) for r in sheets['Material routes'][5:]
                    if r[rh.index('Recipe ID')] == key and r[rh.index('Relationship')] in INGREDIENTS]
            set_cell('Crafting recipes', n, 'Listed ingredient tier', 'Unassigned' if not nums or -1 in nums else f'T{max(nums)}')
        listed[key] = sheets['Crafting recipes'][n-1][ch.index('Listed ingredient tier')]
    h = sheets['Crafting browser'][4]
    for n, row in enumerate(sheets['Crafting browser'][5:], 6):
        get = lambda field: row[h.index(field)]
        if get('Recipe ID') in recipes or get('Input item ID') in selected or get('Output item ID') in selected:
            # A salvage value rank cannot demonstrate a refinement-stage advance.
            interpretation = ('Withheld premise' if get('Route status') == 'Withheld' else
                              'Salvage value tier; crafting readiness separate')
            set_cell('Crafting browser', n, 'Tier interpretation', interpretation)
    audit = dict(date='2026-10-16', authority='Author requested practical-worth brainstorming and exactly 40 of the 200 salvage items in each tier T0 through T4.',
                 method='Relative recoverable utility, replacement labor, material scarcity, condition and plausible specialist demand. These are estimated value bands, not currency prices, drop probabilities or manufacturing stages. Bulk and incomplete matching reduce value. Unknown contents, charge, provenance and unproven powers add no guaranteed value.',
                 tierDefinitions={'T0':'Low-value common goods, incomplete scrap and uncertain reclaiming feed',
                                  'T1':'Modest-value reusable vessels, useful stock and simple fittings',
                                  'T2':'Useful fabricated goods, repair donors and moderate craft labor',
                                  'T3':'High-value prepared materials, skilled tools and artisan mechanisms',
                                  'T4':'Highest relative value: precision assemblies, scarce optics, substantial machinery and costly prepared craft materials'},
                 scope='Local master workbook only. Google Sheets is separate and unsynchronized. Preserve rarity, trade-value fields, item IDs, descriptions, physical condition, recipes, quantities, contents and review/art state. World identity, streaming/unload/reload, generated-object identity, authored placement and persisted runtime deltas are not applicable because no runtime or world objects are changed.',
                 counts={f'T{t}':40 for t in range(5)}, sourceHashBefore=digest(snapshot), items=reviewed)
    OUT.mkdir(parents=True, exist_ok=True)
    assert not AUDIT.exists(), 'Audit already exists; inspect rather than replaying.'
    (OUT / 'before.xlsx').write_bytes(snapshot)
    (OUT / 'plan.json').write_text(json.dumps(dict(audit=audit, changes=changes), indent=2)+'\n')
    # Scoped render inputs retain actual source widths, row heights and text.
    book = openpyxl.load_workbook(BytesIO(snapshot), data_only=True)
    samples = []
    for t in range(5):
        samples.extend([i for i in reviewed if i['tier'] == f'T{t}'][:2])
    preview = []
    for item in samples:
        cells = []
        for col in ['B', 'C']:
            cell = book['Items'][f"{col}{item['row']}"]
            cells.append(dict(value=cell.value, width=book['Items'].column_dimensions[col].width,
                              height=book['Items'].row_dimensions[item['row']].height or 20,
                              font=cell.font.name or 'Arial', size=cell.font.sz or 11,
                              bold=bool(cell.font.b), wrap=bool(cell.alignment.wrap_text)))
        preview.append(dict(cells=cells, tier=item['tier'], reason=item['valueReason']))
    book.close()
    (OUT / 'preview.json').write_text(json.dumps(preview, indent=2))
    print('Prepared 200 individually assessed items, 40 per tier.', {s:len(c) for s,c in changes.items()})


def apply():
    plan = json.loads((OUT / 'plan.json').read_text())
    original = (OUT / 'before.xlsx').read_bytes()
    assert digest(SOURCE.read_bytes()) == plan['audit']['sourceHashBefore'], 'Workbook changed; prepare a fresh assessment.'
    assert not AUDIT.exists()
    authored = openpyxl.load_workbook(OUT / 'authored.xlsx', data_only=True)
    for sheet, cells in plan['changes'].items():
        for address, value in cells.items():
            assert authored[sheet][address].value == value, (sheet, address)
    authored.close()
    result = BytesIO()
    with ZipFile(BytesIO(original)) as archive, ZipFile(result, 'w') as output:
        parts = parts_for(archive)
        replacements = {parts[s]:patch_cells(archive.read(parts[s]).decode(), cells).encode() for s,cells in plan['changes'].items()}
        for member in archive.infolist():
            output.writestr(member, replacements.get(member.filename, archive.read(member.filename)))
    candidate = result.getvalue()
    verify_move(original, candidate, plan['changes'], category_contract=False)
    sheets = parse_sheets(candidate)
    salvage = [r for _,r in master_items(sheets) if r[3] == 'Salvage']
    assert Counter(r[2] for r in salvage) == {f'T{t}':40 for t in range(5)}
    assert [r[0] for _,r in master_items(sheets)] == [r[0] for _,r in master_items(parse_sheets(original))]
    plan['audit'].update(sourceHashAfter=digest(candidate), changedCells={s:len(c) for s,c in plan['changes'].items()},
                         verification='Exact changed-cell allowlist checked across all workbook XML. All original formulas, styles and unrelated cells/archive members retained. All 200 stable IDs retained and each tier has 40 items.')
    (OUT / 'candidate.xlsx').write_bytes(candidate)
    atomic_write(SOURCE, candidate, plan['audit']['sourceHashBefore'])
    atomic_write(AUDIT, (json.dumps(plan['audit'], indent=2)+'\n').encode())
    print('Applied and verified:', plan['audit']['counts'], 'SHA256', digest(candidate))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare', 'apply'])
    prepare() if parser.parse_args().action == 'prepare' else apply()
