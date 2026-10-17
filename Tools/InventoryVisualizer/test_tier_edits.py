"""Exercise tier saves and stale/repeated requests on disposable workbooks."""
import json
import tempfile
import threading
import unittest
from functools import partial
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from uuid import uuid4

import openpyxl
import server
import tier_edits as edits
from test_category_moves import fixture


def tier_fixture(path):
    fixture(path)
    b=openpyxl.load_workbook(path)
    s=b.create_sheet('Taxonomy')
    s.cell(5,1,'Field');s.cell(5,2,'Allowed value')
    for n,t in enumerate(['Not applicable','T0','T1','T2','T3','T4','Unassigned','T5'],6):
        s.cell(n,1,'Material tier');s.cell(n,2,t)
    for name,headers,rows in [
        ('Material routes',['Link ID','Recipe ID','Relationship','Item ID','Material tier','Ingredient tier number','Unresolved ingredient tier'],
         [['input','recipe','Ingredient candidate','test.guard','T1',1,0],['output','recipe','Output','test.tool','T1',-1,0]]),
        ('Crafting recipes',['Recipe ID','Listed ingredient tier'],[['recipe','T1']]),
        ('Tier review',['Item ID','Previous tier','Current tier','Tier sort order'],[['test.guard','Unassigned','T1',1]]),
        ('Craft audit',['Item ID','Tier'],[['test.guard','T1']]),
        ('Crafting browser',['Input item ID','Output item ID','Input tier','Output tier','Recipe ID','Route status','Relationship','Route kind','Tier interpretation'],
         [['test.guard','test.tool','T1','T1','recipe','Proposal','Ingredient candidate','Transformation','Needs tier review']])]:
        if name in b:b.remove(b[name])
        s=b.create_sheet(name)
        for c,h in enumerate(headers,1):s.cell(5,c,h)
        for n,row in enumerate(rows,6):
            for c,v in enumerate(row,1):s.cell(n,c,v)
    b['Tier review']['C6']='=Items!C6'
    b['Craft audit']['B6']='=Items!C6'
    b.save(path);b.close()


class TierTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=server.ROOT/'Logs')
        self.book=Path(self.temp.name)/'fixture.xlsx';self.backups=Path(self.temp.name)/'backups'
        tier_fixture(self.book)

    def tearDown(self):self.temp.cleanup()

    def body(self,**kwargs):
        return dict(itemId='test.guard',fromTier='T1',toTier='T0',sourceHash=edits.digest(self.book.read_bytes()),operationId=str(uuid4()),**kwargs)

    def test_linked_caches_formulas_history_and_idempotency(self):
        original=self.book.read_bytes();body=self.body()
        receipt=edits.update(self.book,self.backups,body)
        self.assertEqual(receipt,edits.update(self.book,self.backups,body))
        sheets=edits.parse_sheets(self.book.read_bytes())
        self.assertEqual(sheets['Items'][5][2],'T0')
        self.assertIn('; t0;',sheets['Items'][5][30])
        self.assertEqual(sheets['Material routes'][5][4:7],('T0',0,0))
        self.assertEqual(sheets['Crafting recipes'][5][1],'T0')
        self.assertEqual(sheets['Crafting browser'][5][2],'T0')
        self.assertEqual(sheets['Crafting browser'][5][-1],'Advances one material tier')
        self.assertEqual(sheets['Tier review'][5][1:4],('Unassigned','T0',0))
        self.assertEqual(sheets['Craft audit'][5][1],'T0')
        self.assertEqual((self.backups/'Tier'/(edits.digest(original)+'.xlsx')).read_bytes(),original)
        w=openpyxl.load_workbook(self.book,data_only=False)
        self.assertEqual(w['Tier review']['C6'].value,'=Items!C6');w.close()
        with self.assertRaises(edits.EditError):edits.update(self.book,self.backups,{**body,'toTier':'T4'})

    def test_non_numeric_tier_recalculates_unresolved_input(self):
        for tier in ['Not applicable','Unassigned']:
            original=self.book.read_bytes()
            result,_=edits.build_tier(original,'test.guard',tier,'T1')
            s=edits.parse_sheets(result)
            self.assertEqual(s['Material routes'][5][5:7],(-1,1))
            self.assertEqual(s['Crafting recipes'][5][1],'Unassigned')
            self.assertEqual(s['Crafting browser'][5][-1],'Tier unresolved')
            self.assertEqual(self.book.read_bytes(),original)

    def test_invalid_stale_noop_and_missing_item(self):
        original=self.book.read_bytes()
        for overrides in [{'toTier':'T6'},{'toTier':None},{'fromTier':'T3'},{'sourceHash':'stale'},{'itemId':'missing'},{'operationId':'../bad'}]:
            with self.assertRaises(edits.EditError):edits.update(self.book,self.backups,{**self.body(),**overrides})
            self.assertEqual(self.book.read_bytes(),original)
        self.assertTrue(edits.update(self.book,self.backups,{**self.body(),'toTier':'T1'})['noop'])
        self.assertEqual(self.book.read_bytes(),original)

    def test_t5_is_numeric_and_saves_linked_ingredient_tiers(self):
        self.assertEqual(edits.tier_number('T5'), 5)
        self.assertEqual(edits.tier_number('T6'), -1)
        self.assertEqual(edits.tier_choices(edits.parse_sheets(self.book.read_bytes())),
                         ['Not applicable','T0','T1','T2','T3','T4','T5','Unassigned'])
        result = edits.update(self.book, self.backups, {**self.body(), 'toTier':'T5'})
        self.assertEqual(result['toTier'], 'T5')
        sheets = edits.parse_sheets(self.book.read_bytes())
        self.assertEqual(sheets['Material routes'][5][4:7], ('T5',5,0))
        self.assertEqual(sheets['Crafting recipes'][5][1], 'T5')
        self.assertEqual(sheets['Tier review'][5][2:4], ('T5',5))
        self.assertEqual(sheets['Craft audit'][5][1], 'T5')

    def test_lost_journal_conflicts_and_locked_workbook(self):
        real=edits.atomic_write;body=self.body()
        def lost(path,*args,**kwargs):
            if path.name.endswith('.committed.json'):raise OSError('journal unavailable')
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=lost):edits.update(self.book,self.backups,body)
        self.assertEqual(edits.operation_status(self.book,self.backups,body['operationId'])['status'],'committed')
        before=self.book.read_bytes()
        next_body={**self.body(),'fromTier':'T0','toTier':'T2'}
        def locked(path,*args,**kwargs):
            if path==self.book:raise PermissionError('Excel lock')
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=locked):
            with self.assertRaises(PermissionError):edits.update(self.book,self.backups,next_body)
        self.assertEqual(self.book.read_bytes(),before)
        self.assertEqual(edits.operation_status(self.book,self.backups,next_body['operationId'])['status'],'not-applied')
        def conflict(path,*args,**kwargs):
            if path==self.book:path.write_bytes(before+b' ')
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=conflict):
            with self.assertRaises(edits.EditError):edits.update(self.book,self.backups,next_body)
        self.assertEqual(self.book.read_bytes(),before+b' ')

    def test_real_format_all_registered_tiers_in_memory(self):
        original=server.WORKBOOK.read_bytes()
        sheets=edits.parse_sheets(original)
        key=next(r[0] for _,r in edits.master_items(sheets) if r[3]=='Crafting components' and r[2]=='T2')
        for tier in edits.tier_choices(sheets):
            result,_=edits.build_tier(original,key,tier,'T2')
            after=edits.parse_sheets(result)
            self.assertEqual(next(r[2] for _,r in edits.master_items(after) if r[0]==key),tier)
        self.assertEqual(server.WORKBOOK.read_bytes(),original)

    def test_http_origin_validation_save_and_catalog(self):
        with patch.object(server,'WORKBOOK',self.book),patch.object(server,'BACKUPS',self.backups):
            http=server.ThreadingHTTPServer(('127.0.0.1',0),partial(server.Handler,directory=str(server.WEB)))
            worker=threading.Thread(target=http.serve_forever,daemon=True);worker.start()
            base=f'http://127.0.0.1:{http.server_port}'
            def post(body,origin=base):
                return urlopen(Request(base+'/api/items/tier',data=json.dumps(body).encode(),headers={'Content-Type':'application/json','Origin':origin,'X-Inventory-Token':server.TOKEN}))
            try:
                with self.assertRaises(HTTPError) as error:post(self.body(),'http://other.test')
                self.assertEqual(error.exception.code,403)
                with self.assertRaises(HTTPError) as error:post({**self.body(),'toTier':'T8'})
                self.assertEqual(error.exception.code,400)
                body=self.body()
                with post(body) as response:self.assertEqual(json.load(response)['toTier'],'T0')
                with urlopen(base+'/api/catalog') as response:
                    catalog=json.load(response);self.assertTrue(catalog['tierEditing']);self.assertEqual(catalog['items'][0]['Tier'],'T0')
                with urlopen(base+'/api/tier-operations/'+body['operationId']) as response:self.assertEqual(json.load(response)['status'],'committed')
            finally:http.shutdown();http.server_close();worker.join()


if __name__=='__main__':unittest.main()
