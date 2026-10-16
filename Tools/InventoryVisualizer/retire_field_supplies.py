"""Move the remaining Field supplies to Crafting components and retire the label."""
import json
import re
from collections import Counter
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import organize_camp_items as guard
import split_equipment as migration
from category_edits import CELL, DEPENDENCIES, parts_for, registry
from workbook_edits import atomic_write, digest, parse_sheets

OUT = migration.ROOT / 'Logs/FieldSuppliesRetirement20261015'
AUDIT = migration.ROOT / 'Docs/Design/Items/FIELD_SUPPLIES_COMPONENT_CLASSIFICATION.json'
OLD, TARGET = 'Field supplies', 'Crafting components'
REASON = 'Author direction October 15: all remaining Field supplies are Crafting components; retire Field supplies as an active category.'


def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    original = migration.SOURCE.read_bytes()
    values = openpyxl.load_workbook(BytesIO(original), data_only=True)
    formulas = openpyxl.load_workbook(BytesIO(original), data_only=False)
    rows = list(values['Items'].values)
    selected = {r[0] for r in rows[5:] if r[0] and r[3] == OLD}
    assert selected and TARGET in registry(list(values['Categories'].values))
    changes = []

    def change(sheet, address, value, cached=False, formula=None):
        changes.append(dict(sheet=sheet, address=address, value=value, cachedOnly=cached,
                            template=None, **({'formula': formula} if formula else {})))

    items = []
    for n, row in enumerate(rows[5:], 6):
        if row[0] not in selected:
            continue
        items.append(dict(id=row[0], name=row[1], previous=OLD))
        change('Items', f'D{n}', TARGET)
        change('Items', f'AF{n}', (row[31] or '') + ' ' + REASON)
        search = list(row)
        search[3] = TARGET
        change('Items', f'AE{n}', '; '.join(str(search[i] or '') for i in
               [0,1,3,4,23,24,25,26,27,28,29,21,2,33,34]).lower(), True)
    for sheet, mappings in DEPENDENCIES.items():
        data = list(values[sheet].values)
        header = data[4]
        for key, category in mappings:
            ki, ci = header.index(key), header.index(category)
            for n, row in enumerate(data[5:], 6):
                if row[ki] in selected:
                    address = f'{migration.col(ci)}{n}'
                    change(sheet, address, TARGET, formulas[sheet][address].data_type == 'f')
    counts = Counter(TARGET if r[0] in selected else r[3] for r in rows[5:] if r[0])
    categories = list(values['Categories'].values)
    retired = 'Retired primary category: all remaining entries are Crafting components. Retained as a historical label only.'
    for n, row in enumerate(categories, 1):
        if len(row) > 4 and row[4] == 'Primary category':
            change('Categories', f'B{n}', counts[row[0]], formula=formulas['Categories'][f'B{n}'].value[1:])
        if row[0] == OLD:
            change('Categories', f'C{n}', retired)
            change('Categories', f'D{n}', REASON)
            change('Categories', f'E{n}', 'Legacy category')
    for n, row in enumerate(values['Taxonomy'].values, 1):
        if row[:2] == ('Category', OLD):
            change('Taxonomy', f'C{n}', retired)
    choices = sorted(c for c in registry(categories) if c != OLD)
    for n, choice in enumerate(choices + [' '], 6):
        change('Categories', f'G{n}', choice)
    plan = dict(sourceHash=digest(original), target=TARGET, items=items, changes=changes,
                counts=dict(counts), extendTables={},
                previousChoices=formulas.defined_names['InventoryCategoryChoices'].attr_text,
                categoryChoices=f"'Categories'!$G$6:$G${len(choices)+5}")
    (OUT / 'before.xlsx').write_bytes(original)
    (OUT / 'plan.json').write_text(json.dumps(plan, indent=2), encoding='utf-8')
    values.close()
    formulas.close()
    print('Prepared', len(items), 'component classifications and category retirement.')


def apply():
    plan = json.loads((OUT / 'plan.json').read_text(encoding='utf-8'))
    source = migration.SOURCE
    candidate_path = OUT / 'candidate.xlsx'
    candidate_path.write_bytes((OUT / 'before.xlsx').read_bytes())
    migration.SOURCE, migration.OUT = candidate_path, OUT
    try:
        migration.apply()
    finally:
        migration.SOURCE = source
    staged = candidate_path.read_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(staged)) as archive, ZipFile(output, 'w') as result:
        category_part = parts_for(archive)['Categories']
        last_choice = re.search(r'\$(\d+)$', plan['previousChoices']).group(1)
        for info in archive.infolist():
            raw = archive.read(info.filename)
            if info.filename == 'xl/workbook.xml':
                old, new = plan['previousChoices'].encode(), plan['categoryChoices'].encode()
                assert raw.count(old) == 1
                raw = raw.replace(old, new)
            if info.filename == category_part:
                raw = re.sub(CELL, lambda m: '' if m.group(1) == f'G{last_choice}' else m.group(),
                             raw.decode(), flags=re.S).encode()
            result.writestr(info, raw)
    candidate = output.getvalue()
    sheets = parse_sheets(candidate)
    assert OLD not in registry(sheets['Categories'])
    assert not any(r[3] == OLD for r in sheets['Items'][5:] if r[0])
    assert sum(r[3] == TARGET for r in sheets['Items'][5:] if r[0]) == plan['counts'][TARGET]
    assert [r[0] for r in sheets['Items'][5:] if r[0]] == [r[0] for r in parse_sheets((OUT/'before.xlsx').read_bytes())['Items'][5:] if r[0]]
    check = openpyxl.load_workbook(BytesIO(candidate), read_only=True)
    assert check.defined_names['InventoryCategoryChoices'].attr_text == plan['categoryChoices']
    check.close()
    assert not AUDIT.exists()
    guard.release_guard()
    atomic_write(source, candidate, plan['sourceHash'])
    audit = dict(reason=REASON, items=plan['items'], retiredCategory=OLD,
                 sourceHashBefore=plan['sourceHash'], sourceHashAfter=digest(candidate),
                 scope='Catalog classification only; IDs, properties, recipes, review/art state, world identity, fixed terrain, streaming and persisted runtime deltas unchanged.')
    atomic_write(AUDIT, (json.dumps(audit, indent=2) + '\n').encode())
    (OUT / 'after.xlsx').write_bytes(candidate)
    print('Verified', len(plan['items']), 'moves; Field supplies retired.')


if __name__ == '__main__':
    guard.OUT, guard.prepare, guard.apply = OUT, prepare, apply
    guard.run()
