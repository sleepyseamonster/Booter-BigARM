"""Register T5 and move Ancient salvage there with a preserving cell merge."""
import argparse
import json
import re
from collections import Counter
from html import unescape
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter
from balance_catalog import merge
from category_edits import CELL, parts_for, verify_move
from tier_edits import tier_number
from workbook_edits import atomic_write, digest, master_items, parse_sheets
import organize_camp_items as guard

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT/'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
OUT = ROOT/'Logs/TierFive20261016'
AUDIT = ROOT/'Docs/Design/Items/TIER_FIVE_ANCIENT_SALVAGE_2026-10-16.json'
DEFINITION = ('T0-T4 retain their current practical-worth assignments. T5 is the author-selected '
              'ancient salvage band, initially containing only the Ancient salvage category. '
              'Tier is independent of rarity, condition, crafting readiness and manufacturing capability. '
              'Not applicable and Unassigned remain available as legacy options.')


def metadata(snapshot, plan):
    """Only extend the two registry tables, new-row styles and tier validations."""
    stream = BytesIO()
    with ZipFile(BytesIO(snapshot)) as source, ZipFile(stream,'w') as target:
        parts = parts_for(source)
        replacements = {}
        for sheet, end in plan['ends'].items():
            raw = source.read(parts[sheet]).decode()
            raw = re.sub(r'(<(?:\w+:)?dimension\b[^>]*\bref="[^"]*?)(\d+)(")',
                         lambda m:m[1]+str(end)+m[3],raw,count=1)
            # New registry records inherit their adjacent family record styles.
            template = 226 if sheet=='Taxonomy' else 47
            cells = {m[1]:m[0] for m in re.finditer(CELL,raw,re.S)}
            for column in ('ABC' if sheet=='Taxonomy' else 'ABCDE'):
                old = cells.get(f'{column}{template}','')
                new = cells.get(f'{column}{end}')
                style = re.search(r'\bs="([^"]+)"',old)
                if new and style and not re.search(r'\bs="',new):
                    revised = re.sub(r'(<(?:\w+:)?c\b)',lambda m:m[1]+f' s="{style[1]}"',new,count=1)
                    raw = raw.replace(new,revised,1)
            if sheet=='Categories':
                address = f'B{end}'
                cell = next((m[0] for m in re.finditer(CELL,raw,re.S) if m[1]==address),None)
                if cell and not re.search(r'<(?:\w+:)?f\b',cell):
                    prefix = re.search(r'<((?:\w+:)?)c\b',cell)[1]
                    revised = cell.replace(f'<{prefix}v>',f'<{prefix}f>COUNTIF(Items!$C$6:$C$1005,"T5")</{prefix}f><{prefix}v>',1)
                    raw = raw.replace(cell,revised,1)
            replacements[parts[sheet]] = raw.encode()
        raw = source.read(parts['Items']).decode()
        choices = '"T0,T1,T2,T3,T4,T5,Unassigned,Not applicable"'
        def validation(match):
            node = match[0]
            if re.search(r'\bsqref="(?:C6:C1005|AH6:AH1005)"',node):
                node = re.sub(r'(<(?:\w+:)?formula1>).*?(</(?:\w+:)?formula1>)',
                              lambda m:m[1]+choices+m[2],node,flags=re.S)
            return node
        replacements[parts['Items']] = re.sub(r'<(?:\w+:)?dataValidation\b[^>]*>.*?</(?:\w+:)?dataValidation>',validation,raw,flags=re.S).encode()
        for member in source.namelist():
            if member.startswith('xl/tables/') and member.endswith('.xml'):
                raw = source.read(member).decode()
                for table, sheet in [('ItemTaxonomy','Taxonomy'),('CategoriesTable','Categories')]:
                    if f'name="{table}"' in raw:
                        raw = re.sub(r'(\bref="[A-Z]+\d+:[A-Z]+)\d+"',lambda m:m[1]+str(plan['ends'][sheet])+'"',raw)
                        replacements[member] = raw.encode()
        for info in source.infolist():
            target.writestr(info,replacements.get(info.filename,source.read(info)))
    return stream.getvalue()


def build(snapshot):
    sheets = parse_sheets(snapshot)
    assert not any(row and row[:2]==('Material tier','T5') for row in sheets['Taxonomy']), 'T5 is already registered; inspect rather than replay.'
    selected = [(n,dict(zip(sheets['Items'][4],row))) for n,row in master_items(sheets) if row[3]=='Ancient salvage']
    assert selected, 'Ancient salvage category is empty.'
    assert not any(row[2]=='T5' for _,row in master_items(sheets)), 'Review existing T5 assignments before initialization.'
    ids = {row['Item ID'] for _,row in selected}
    changes, formulas = {}, {}
    original = parse_sheets(snapshot)
    def set_cell(sheet, number, field, value):
        h = sheets[sheet][4]
        if field not in h:return
        column = h.index(field)
        row = list(sheets[sheet][number-1])
        if row[column]!=value:
            changes.setdefault(sheet,{})[f'{get_column_letter(column+1)}{number}'] = value
            row[column] = value
            sheets[sheet][number-1] = tuple(row)
    for number,item in selected:
        set_cell('Items',number,'Tier','T5')
        set_cell('Items',number,'Tier basis','Author direction October 16, 2026: T5 ancient salvage band. Ancient origin and replacement difficulty justify the separate tier. Rarity, condition, compatibility and crafting qualification remain independent; no reproduction capability is granted.')
        set_cell('Items',number,'Tier sort order',5)
        row = dict(zip(sheets['Items'][4],sheets['Items'][number-1]))
        fields = ['Item ID','Item name','Category','Subtype','Item type','Processing stage','Origin family','Primary use','Material family','Physical form','Design status','Search aliases','Tier','Finished-item input tier','Crafting role']
        set_cell('Items',number,'Search terms','; '.join(str(row.get(k) or '') for k in fields).lower())
        set_cell('Items',number,'Crafting coverage','Route listed; review status/completeness' if (row.get('Listed output routes') or 0)>0 else 'No production route listed')
    for sheet,mappings in {'Tier review':[('Item ID','Current tier'),('Item ID','Tier sort order')], 'Craft audit':[('Item ID','Tier')]}.items():
        h = sheets[sheet][4]
        for n,row in enumerate(sheets[sheet][5:],6):
            for key,field in mappings:
                if key in h and field in h and row[h.index(key)] in ids:
                    set_cell(sheet,n,field,5 if field=='Tier sort order' else 'T5')
    # Ancient recovery has no player-manufacture routes. Do not invent one.
    for sheet in ['Material routes','Crafting browser']:
        assert not any(any(v in ids for v in row if isinstance(v,str)) for row in sheets[sheet][5:]), 'Ancient route topology needs explicit review.'
    for n,row in enumerate(sheets['Taxonomy'],1):
        if row and row[0]=='Material tier':
            set_cell('Taxonomy',n,'Definition / rule',DEFINITION)
    ends = {'Taxonomy':len(sheets['Taxonomy'])+1,'Categories':len(sheets['Categories'])+1}
    changes.setdefault('Taxonomy',{}).update({f'{c}{ends["Taxonomy"]}':v for c,v in zip('ABC',['Material tier','T5',DEFINITION])})
    changes.setdefault('Categories',{}).update({f'{c}{ends["Categories"]}':v for c,v in zip('ABCDE',['T5 ancient salvage',len(selected),'Author-selected Tier 5 band, initially assigned only to Ancient salvage.','Rarity and qualification remain independent.','Item family'])})
    guide = {
        'Tier versus rarity and quality':DEFINITION,
        'Tier fields':'Items.C is the current value tier T0-T5; AG records its basis. T5 initially contains the Ancient salvage category only. Previous tier and finished-input fields retain historical planning information. T5 grants no manufacturing or crafting readiness.',
        'Sort and filter by tier':'Filter Items.C by T0 through T5. Sort Tier sort order in AN for numeric tiers first, then Unassigned (98) and Not applicable (99). T5 is initially Ancient salvage only. Filter and sort whole native tables. Tier and rarity are independent.',
        'Tier 4 boundaries and remaining work':'Historical October 15 T4 planning remains separate from current value tiers. October 16 author direction adds T5 for Ancient salvage; it does not extend the contemporary manufacturing ladder or establish ancient reproduction. Quantities, yields, tools, power, tolerances and qualification remain open.'
    }
    for n,row in enumerate(sheets['Guide'],1):
        if row and row[0] in guide:set_cell('Guide',n,'Guidance',guide[row[0]])
    # Extend existing numeric helpers so Excel recalculation retains T5 support.
    with ZipFile(BytesIO(snapshot)) as archive:
        parts = parts_for(archive)
        for sheet,column in [('Items','AN'),('Material routes','P')]:
            for match in re.finditer(CELL,archive.read(parts[sheet]).decode(),re.S):
                address = match[1]
                if not re.fullmatch(column+r'\d+',address):continue
                tag = re.search(r'<(?:\w+:)?f\b[^>]*>(.*?)</(?:\w+:)?f>',match[0],re.S)
                if not tag:continue
                expression = unescape(tag[1]);n = int(re.search(r'\d+',address)[0])
                if sheet=='Items':
                    revised = expression.replace(f'IF(C{n}="T4",4,',f'IF(C{n}="T4",4,IF(C{n}="T5",5,')+')'
                    assert revised!=expression+')', (sheet,address,'unrecognized sort formula')
                    tier = sheets[sheet][n-1][2]
                    value = tier_number(tier) if tier_number(tier)>=0 else 98 if tier=='Unassigned' else 99
                else:
                    revised = expression.replace(f'IF(F{n}="T4",4,-1)',f'IF(F{n}="T4",4,IF(F{n}="T5",5,-1))')
                    assert revised!=expression,(sheet,address,'unrecognized ingredient formula')
                    value = original[sheet][n-1][15]
                    assert value is not None,(sheet,address,'missing original ingredient cache')
                formulas.setdefault(sheet,{})[address] = revised
                changes.setdefault(sheet,{})[address] = value
    plan = dict(changes=changes,formulas=formulas,ends=ends)
    candidate = metadata(merge(snapshot,changes,formulas),plan)
    # Normalize only declared metadata/formula edits, then prove every other
    # archive member, cell, formula and existing style is unchanged.
    previous = {sheet:{address:original[sheet][int(re.search(r'\d+',address)[0])-1][openpyxl.utils.cell.column_index_from_string(re.match(r'[A-Z]+',address)[0])-1] or 0 for address in expressions} for sheet,expressions in formulas.items()}
    normalized = metadata(merge(snapshot,previous,formulas),plan)
    verify_move(normalized,candidate,changes,category_contract=False)
    before = {row[0]:row for _,row in master_items(original)}
    after_sheets = parse_sheets(candidate)
    after = {row[0]:row for _,row in master_items(after_sheets)}
    assert list(before)==list(after)
    for key,row in before.items():
        assert row[15]==after[key][15], 'Rarity changed'
        if key not in ids:
            assert row[:39]==after[key][:39], 'Unrelated item changed'
        else:
            assert after[key][2]=='T5' and after[key][39]==5
    assert {row[0] for row in after.values() if row[2]=='T5'}==ids
    receipt = dict(date='2026-10-16',sourceHash=digest(snapshot),resultHash=digest(candidate),count=len(ids),
        rarities=dict(Counter(item['Rarity'] for _,item in selected)),
        items=[dict(id=item['Item ID'],name=item['Item name'],previousTier=item['Tier'],tier='T5',rarity=item['Rarity']) for _,item in selected],
        changedCells={sheet:len(cells) for sheet,cells in changes.items()},formulaUpdates={sheet:len(cells) for sheet,cells in formulas.items()},
        scope='Catalog tier registration and Ancient salvage assignment only. Preserve IDs, rarities, names, descriptions, condition, recipe topology, historical tiers and review/art state. World identity, streaming, generated-object identity, authored placement and persisted runtime deltas are not applicable: no runtime system or object changes.')
    plan['receipt'] = receipt
    return candidate,plan


def prepare():
    OUT.mkdir(parents=True,exist_ok=True)
    snapshot = SOURCE.read_bytes()
    candidate,plan = build(snapshot)
    (OUT/'before.xlsx').write_bytes(snapshot)
    (OUT/'candidate.xlsx').write_bytes(candidate)
    (OUT/'plan.json').write_text(json.dumps(plan,indent=2),encoding='utf-8')
    old = {r[0]:r for _,r in master_items(parse_sheets(snapshot))}
    samples = []
    for previous in dict.fromkeys(item['previousTier'] for item in plan['receipt']['items']):
        for item in [i for i in plan['receipt']['items'] if i['previousTier']==previous][:2]:
            row = old[item['id']]
            samples.append(dict(name=row[1],previous=row[2],tier='T5',rarity=row[15]))
    (OUT/'preview.json').write_text(json.dumps(samples),encoding='utf-8')
    print('Prepared T5 registration for',plan['receipt']['count'],'ancient items.',flush=True)


def apply():
    plan = json.loads((OUT/'plan.json').read_text())
    book = openpyxl.load_workbook(OUT/'authored.xlsx',data_only=True)
    for sheet,cells in plan['changes'].items():
        for address,value in cells.items():assert book[sheet][address].value==value,(sheet,address,'authored value')
    expected = [[sheet,address,expression] for sheet,cells in plan['formulas'].items() for address,expression in cells.items()]
    assert [list(row) for row in list(book['Formula edits'].values)[1:]]==expected
    book.close()
    guard.release_guard()
    atomic_write(SOURCE,(OUT/'candidate.xlsx').read_bytes(),plan['receipt']['sourceHash'])
    atomic_write(AUDIT,(json.dumps(plan['receipt'],indent=2)+'\n').encode())
    print('Saved T5 for',plan['receipt']['count'],'Ancient salvage items; rarities unchanged.',flush=True)


if __name__=='__main__':
    parser = argparse.ArgumentParser();parser.add_argument('step',choices=['prepare','apply','run'])
    step = parser.parse_args().step
    if step=='run':
        guard.OUT=OUT;guard.prepare=prepare;guard.apply=apply;guard.run()
    else:globals()[step]()
