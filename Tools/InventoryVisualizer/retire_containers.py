"""Reclassify the remaining container proposals as ordinary salvage."""
import argparse
import json
import re
from collections import Counter
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import organize_camp_items as guard
import split_equipment as migration
from category_edits import CELL, DEPENDENCIES, parts_for
from workbook_edits import atomic_write, digest

OUT = migration.ROOT / 'Logs/ContainersToSalvage20261015'
AUDIT = migration.ROOT / 'Docs/Design/Items/CONTAINERS_SALVAGE_CLASSIFICATION.json'
REASON = 'Author direction October 15: remaining Containers entries are ordinary salvage for possible reuse or material recovery. Specialized container, payload, transfer, sealing and cargo mechanics are withdrawn; no specific reuse or recovery yield is approved.'
PURPOSE = 'Ordinary found salvage; possible reuse or material recovery after inspection.'
LIMIT = 'No special container behavior, capacity, payload, sealing, sample handling or carrier eligibility. Specific reuse and recovery yields remain undecided.'


def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    original = migration.SOURCE.read_bytes()
    values = openpyxl.load_workbook(BytesIO(original), data_only=True)
    formulas = openpyxl.load_workbook(BytesIO(original), data_only=False)
    rows = list(values['Items'].values)
    selected = {r[0]: r for r in rows[5:] if r[0] and r[3] == 'Containers'}
    assert selected, 'No remaining Containers entries'
    changes = {}

    def change(sheet, address, value, cached=False, formula=None):
        changes[sheet, address] = dict(sheet=sheet, address=address, value=value,
            cachedOnly=cached, template=None, **({'formula': formula} if formula else {}))

    for n, row in enumerate(rows[5:], 6):
        if row[0] not in selected:
            continue
        updated = list(row)
        replacements = {3: 'Salvage', 4: 'Reusable salvage', 6: PURPOSE,
            8: 'Ordinary salvage inventory / cargo; no special carrier rules',
            11: 'Found salvage; location and availability unresolved',
            12: 'No dedicated container production route; possible salvage reuse remains undecided.',
            13: 'No special container behavior', 14: LIMIT,
            20: 'Possible reuse, material recovery and yields remain open. Former specialized container proposal withdrawn.',
            23: 'Salvage', 25: 'Recovered / found', 26: 'Salvage reuse',
            31: (row[31] or '') + ' ' + REASON,
            34: 'Salvage / possible reuse', 38: 'Found salvage; former container production proposal retired'}
        for ci, value in replacements.items():
            updated[ci] = value
            change('Items', f'{migration.col(ci)}{n}', value)
        search = '; '.join(str(updated[i] or '') for i in [0,1,3,4,23,24,25,26,27,28,29,21,2,33,34]).lower()
        change('Items', f'AE{n}', search, True)

    for sheet, mappings in DEPENDENCIES.items():
        data = list(values[sheet].values)
        headers = data[4]
        for key, category in mappings:
            ki, ci = headers.index(key), headers.index(category)
            for n, row in enumerate(data[5:], 6):
                if row[ki] in selected:
                    address = f'{migration.col(ci)}{n}'
                    change(sheet, address, 'Salvage', formulas[sheet][address].data_type == 'f')

    # Retain linked records and identities as history; explicitly withdraw their old roles.
    retired_routes = set()
    for row in list(values['Material routes'].values)[5:]:
        if row[3] in selected and str(row[2]).lower().startswith('output'):
            retired_routes.add(row[1])
    for sheet in ['Item uses', 'Container concepts', 'Craft audit', 'Material routes', 'Crafting browser', 'Crafting recipes']:
        data = list(values[sheet].values)
        for n, row in enumerate(data[5:], 6):
            if sheet == 'Item uses' and row[0] in selected:
                for ci, value in {4:'Salvage reuse', 5:REASON}.items():
                    change(sheet, f'{migration.col(ci)}{n}', value)
            elif sheet == 'Container concepts' and row[0] in selected:
                for ci, value in {2:'Ordinary salvage', 4:PURPOSE, 5:'Unspecified', 6:'No special closure mechanic', 7:'No assigned contents or payload', 8:LIMIT, 9:REASON}.items():
                    change(sheet, f'{migration.col(ci)}{n}', value)
            elif sheet == 'Craft audit' and row[0] in selected:
                for ci, value in {2:'Salvage', 4:'Found salvage', 6:'No production route required', 7:'No dedicated station or tools', 8:'Former container production proposal retired', 9:REASON}.items():
                    change(sheet, f'{migration.col(ci)}{n}', value)
            elif sheet == 'Material routes':
                if row[3] in selected:
                    change(sheet, f'H{n}', 'Salvage')
                if row[1] in retired_routes:
                    for ci,value in {10:'Retired',12:'Withdrawn container proposal',13:REASON}.items():
                        change(sheet, f'{migration.col(ci)}{n}', value)
            elif sheet == 'Crafting browser':
                for ki,ci in [(18,2),(19,7)]:
                    if row[ki] in selected:
                        change(sheet, f'{migration.col(ci)}{n}', 'Salvage')
                if row[9] in retired_routes:
                    change(sheet, f'L{n}', 'Retired')
                    change(sheet, f'M{n}', 'Withdrawn container proposal')
            elif sheet == 'Crafting recipes' and row[0] in retired_routes:
                for ci,value in {3:'Retired',4:'Withdrawn container proposal',5:REASON,6:'No dedicated station required for found salvage',12:'Possible salvage reuse remains undecided'}.items():
                    change(sheet, f'{migration.col(ci)}{n}', value)

    counts = Counter('Salvage' if r[0] in selected else r[3] for r in rows[5:] if r[0])
    categories = list(values['Categories'].values)
    retired = 'Retired primary category: all remaining entries moved to Salvage by author direction. No special container mechanics; ordinary salvage may be reused in other ways.'
    for n,row in enumerate(categories, 1):
        if len(row)>4 and row[4]=='Primary category':
            change('Categories', f'B{n}', counts[row[0]], formula=formulas['Categories'][f'B{n}'].value[1:])
        if row[0]=='Containers':
            for ci,value in {2:retired,3:REASON,4:'Legacy category'}.items():
                change('Categories',f'{migration.col(ci)}{n}',value)
        if row[0]=='Salvage':
            change('Categories',f'C{n}',(row[2] or '')+' Includes former container objects as ordinary salvage without special container mechanics.')
    for n,row in enumerate(values['Taxonomy'].values,1):
        if row[:2]==('Category','Containers'):
            change('Taxonomy',f'C{n}',retired)
    choices = sorted(r[0] for r in categories if len(r)>4 and r[4]=='Primary category' and r[0]!='Containers')
    for n,value in enumerate(choices+[' '],6):
        change('Categories',f'G{n}',value)
    plan = dict(sourceHash=digest(original),target='Salvage',
        items=[dict(id=r[0],name=r[1],previous=r[3],previousSubtype=r[4]) for r in selected.values()],
        changes=list(changes.values()),counts=dict(counts),extendTables={},retiredRoutes=sorted(retired_routes),
        previousChoices=formulas.defined_names['InventoryCategoryChoices'].attr_text,
        categoryChoices=f"'Categories'!$G$6:$G${len(choices)+5}")
    (OUT/'before.xlsx').write_bytes(original)
    (OUT/'plan.json').write_text(json.dumps(plan,indent=2),encoding='utf-8')
    values.close(); formulas.close()
    print('Prepared',len(selected),'salvage entries and retired',len(retired_routes),'container production proposals')


def apply():
    plan=json.loads((OUT/'plan.json').read_text(encoding='utf-8'))
    source=migration.SOURCE
    candidate_path=OUT/'candidate.xlsx'
    candidate_path.write_bytes((OUT/'before.xlsx').read_bytes())
    migration.SOURCE=candidate_path; migration.OUT=OUT
    try:
        migration.apply()
    finally:
        migration.SOURCE=source
    staged=candidate_path.read_bytes()
    # Shrink the existing native validation name after retiring the empty category.
    output=BytesIO()
    with ZipFile(BytesIO(staged)) as archive, ZipFile(output,'w') as result:
        category_part=parts_for(archive)['Categories']
        last_choice=re.search(r'\$(\d+)$',plan['previousChoices']).group(1)
        for info in archive.infolist():
            raw=archive.read(info.filename)
            if info.filename=='xl/workbook.xml':
                old=plan['previousChoices'].encode(); new=plan['categoryChoices'].encode()
                assert raw.count(old)==1
                raw=raw.replace(old,new)
            if info.filename==category_part:
                raw=re.sub(CELL,lambda m: '' if m.group(1)==f'G{last_choice}' else m.group(),raw.decode(),flags=re.S).encode()
            result.writestr(info,raw)
    candidate=output.getvalue()
    check=openpyxl.load_workbook(BytesIO(candidate),data_only=True)
    assert check.defined_names['InventoryCategoryChoices'].attr_text==plan['categoryChoices']
    assert not any(r[3]=='Containers' for r in list(check['Items'].values)[5:] if r[0])
    assert all(check['Categories'][f'B{n}'].value==plan['counts'].get(r[0],0)
        for n,r in enumerate(check['Categories'].values,1) if len(r)>4 and r[4]=='Primary category')
    check.close()
    candidate_path.write_bytes(candidate)
    assert not AUDIT.exists()
    guard.release_guard()
    atomic_write(source,candidate,plan['sourceHash'])
    audit=dict(reason=REASON,items=plan['items'],retiredRoutes=plan['retiredRoutes'],
        sourceHashBefore=plan['sourceHash'],sourceHashAfter=digest(candidate),
        scope='Catalog and proposal records only. Stable IDs, names, tiers, physical materials/forms and review/art links retained. No runtime world identity, streaming, placement or persisted gameplay delta changes.')
    atomic_write(AUDIT,(json.dumps(audit,indent=2)+'\n').encode())
    print('Saved',len(plan['items']),'ordinary salvage entries; Containers retired from category choices.')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['prepare','apply','run'])
    action=parser.parse_args().action
    if action=='run':
        guard.OUT=OUT;guard.prepare=prepare;guard.apply=apply;guard.run()
    else:
        prepare() if action=='prepare' else apply()
