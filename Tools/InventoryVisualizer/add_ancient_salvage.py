"""Add the three author-selected salvage concepts through a preserving OOXML merge.

Whole-workbook Artifact Tool import previously exhausted its heap. This uses the
established narrow-edit fallback, preserving unrelated workbook parts and cells.
"""
import json
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

import openpyxl
from openpyxl.formula.translate import Translator
from category_edits import CELL, NS, parts_for, patch_cells
from workbook_edits import atomic_write, digest

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
DATA = ROOT / 'Docs/Design/Items/ANCIENT_SALVAGE_PROPOSALS.json'
OUT = ROOT / 'Logs/AncientSalvageThree'


def build(snapshot):
    spec = json.loads(DATA.read_text(encoding='utf-8'))
    w = openpyxl.load_workbook(BytesIO(snapshot), data_only=False)
    v = openpyxl.load_workbook(BytesIO(snapshot), data_only=True)
    ids = [r[0] for r in list(v['Items'].values)[5:] if r[0]]
    assert len(ids) == len(set(ids)), 'Duplicate master IDs'
    assert not any(i['id'] in ids for i in spec['items']), 'Proposal already exists'
    assert not any(r[1] in {i['name'] for i in spec['items']} for r in list(v['Items'].values)[5:] if r[0]), 'Name already exists'
    last = int(re.search(r'\d+$', w['Items'].tables['ItemsTable'].ref).group())
    assert last + 3 <= 1005, 'Master formula capacity exceeded'
    changed, new_rows, tables = {}, {}, {}
    summary = {}
    with ZipFile(BytesIO(snapshot)) as archive:
        parts = parts_for(archive)
        def row_xml(sheet, template, number, values, caches):
            raw = archive.read(parts[sheet]).decode('utf-8')
            prefix = re.search(r'<((?:\w+:)?)worksheet\b', raw).group(1)
            source_row = re.search(rf'<(?:\w+:)?row\b[^>]*\br="{template}"[^>]*>.*?</(?:\w+:)?row>', raw, re.S).group()
            styles = {re.sub(r'\d+', '', m.group(1)): m.group() for m in re.finditer(CELL, source_row, re.S)}
            result = []
            for index, value in enumerate(values, 1):
                col = openpyxl.utils.get_column_letter(index);address = f'{col}{number}'
                cell = styles.get(col, f'<{prefix}c r="{col}{template}"/>')
                opening = cell[:cell.index('>')].rstrip('/')
                opening = re.sub(r'\br="[A-Z]+\d+"', f'r="{address}"', opening)
                opening = re.sub(r'\s+t="[^"]*"', '', opening)
                if value is None:
                    result.append(opening + '/>')
                elif isinstance(value, str) and value.startswith('='):
                    cache = caches[index]
                    kind = '' if isinstance(cache, (int, float)) else ' t="str"'
                    result.append(opening + kind + f'><{prefix}f>{escape(value[1:])}</{prefix}f><{prefix}v>{escape(str(cache))}</{prefix}v></{prefix}c>')
                elif isinstance(value, (int, float)):
                    result.append(opening + f'><{prefix}v>{value}</{prefix}v></{prefix}c>')
                else:
                    result.append(opening + f' t="inlineStr"><{prefix}is><{prefix}t xml:space="preserve">{escape(value)}</{prefix}t></{prefix}is></{prefix}c>')
            new_rows.setdefault(sheet, {})[number] = (values, caches)
            return f'<{prefix}row r="{number}">' + ''.join(result) + f'</{prefix}row>'
        def insert(sheet, rows):
            raw = archive.read(parts[sheet]).decode('utf-8')
            prefix = re.search(r'<((?:\w+:)?)worksheet\b', raw).group(1)
            for number, content in rows:
                assert not any(c.value is not None for c in w[sheet][number]), (sheet, number, 'occupied row')
                pattern = rf'<(?:\w+:)?row\b[^>]*\br="{number}"[^>]*(?:/>|>.*?</(?:\w+:)?row>)'
                if re.search(pattern, raw, re.S):raw = re.sub(pattern, lambda m: content, raw, count=1, flags=re.S)
                else:raw = raw.replace(f'</{prefix}sheetData>', content + f'</{prefix}sheetData>')
            dimension = re.search(r'(<(?:\w+:)?dimension\b[^>]*\bref=")([^"]+)', raw)
            if dimension:
                old_ref = dimension.group(2)
                end = max(int(re.search(r'\d+$', old_ref).group()), max(n for n, _ in rows))
                new_ref = re.sub(r'\d+$', str(end), old_ref)
                raw = raw[:dimension.start(2)] + new_ref + raw[dimension.end(2):]
            changed[parts[sheet]] = raw.encode('utf-8')
        items, uses, audits = [], [], []
        use_last = w['Item uses'].max_row
        for offset, item in enumerate(spec['items'], 1):
            n = last + offset
            fields = [w['Items'].cell(631, c).value for c in range(1, 41)]
            for c, value in enumerate(fields, 1):
                if isinstance(value, str) and value.startswith('='):
                    fields[c-1] = Translator(value, origin=f'{openpyxl.utils.get_column_letter(c)}631').translate_formula(f'{openpyxl.utils.get_column_letter(c)}{n}')
            overrides = {0:item['id'],1:item['name'],2:'T0',3:'Ancient salvage',4:'Ancient recovery',5:'New proposal',6:item['purpose'],
                         11:'Recover from buried ancient machinery; specific sites and recovery methods remain proposals.',
                         12:'Inspect and preserve the original. No manufacturing or internal repair recipe established.',
                         14:'Extremely rare. Condition, compatibility, operating requirements and remaining useful life are unresolved.',
                         15:'Extremely rare',20:'Artwork direction: '+item['icon']+' Condition and compatibility determine usefulness; original manufacture remains beyond established contemporary crafting.',
                         21:item['name'].lower()+'; ancient salvage; as-found recovery',22:spec['source']+' Setting constraints: Docs/WORLD_BASIS.md#material-culture.',
                         23:'Machine part',24:'Recovered part',25:'Ancient technology',26:item['use'],27:item['material'],28:'Assembly',29:'Proposal',
                         31:'Author-selected brainstorm proposal; catalog creation does not establish functioning technology or a recipe.',
                         32:'T0 is the as-found recovery stage. Inspection and qualification of a compatible original would require a separate readiness record; rarity is independent of tier.',
                         33:'T0',34:'Component'}
            for c, value in overrides.items():fields[c] = value
            search = '; '.join(str(fields[c] or '') for c in [0,1,3,4,23,24,25,26,27,28,29,21,2,33,34]).lower()
            caches = {31:search,36:0,37:0,38:0,39:'Raw entry feed; no production route',40:0}
            items.append((n,row_xml('Items',631,n,fields,caches)))
            u = use_last + offset
            use_fields = [w['Item uses'].cell(use_last,c).value for c in range(1,7)]
            for c,value in enumerate(use_fields,1):
                if isinstance(value,str) and value.startswith('='):use_fields[c-1]=Translator(value,origin=f'{openpyxl.utils.get_column_letter(c)}{use_last}').translate_formula(f'{openpyxl.utils.get_column_letter(c)}{u}')
            use_fields[0]=item['id'];use_fields[5]='Author-selected salvage proposal; functional use awaits qualification.'
            uses.append((u,row_xml('Item uses',use_last,u,use_fields,{2:item['name'],3:'Ancient salvage',5:item['use']})))
            audit=[item['id'],f'=Items!B{n}',f'=Items!X{n}',f'=Items!C{n}','As-found ancient recovery',f'=Items!AJ{n}','No reproduction route established','Inspection equipment and compatibility criteria remain undefined.','Condition unknown; qualification required',fields[20]]
            audits.append((n,row_xml('Craft audit',631,n,audit,{2:item['name'],3:'Machine part',4:'T0',6:0})))
        insert('Items',items);insert('Item uses',uses);insert('Craft audit',audits)
        for sheet, table, end in [('Items','ItemsTable',last+3),('Item uses','ItemUseRelationships',use_last+3),('Craft audit','CraftAudit',last+3)]:
            oldref=w[sheet].tables[table].ref;newref=re.sub(r'\d+$',str(end),oldref)
            for member in archive.namelist():
                if member.startswith('xl/tables/') and member.endswith('.xml'):
                    raw=archive.read(member).decode('utf-8')
                    if f'name="{table}"' in raw:
                        changed[member]=raw.replace(f'ref="{oldref}"',f'ref="{newref}"').encode('utf-8');tables[table]=newref
        counts=Counter(r[3] for r in list(v['Items'].values)[5:] if r[0]);counts['Ancient salvage']+=3
        for r in w['Categories']:
            if r[0].value=='Ancient salvage':summary.setdefault('Categories',{})[r[1].coordinate]=counts['Ancient salvage']
            if r[0].value=='Ancient recovered items':summary.setdefault('Categories',{})[r[1].coordinate]=sum(1 for row in list(v['Items'].values)[5:] if row[0] and row[25]=='Ancient technology')+3
        for r in w['Collection targets']:
            if r[0].value=='Machine parts':
                count=sum(1 for row in list(v['Items'].values)[5:] if row[0] and row[23]=='Machine part' and row[29]!='Withheld assumption')+3
                summary['Collection targets']={r[3].coordinate:count,r[4].coordinate:max(0,r[2].value-count)}
        for sheet, cells in summary.items():changed[parts[sheet]]=patch_cells(archive.read(parts[sheet]).decode('utf-8'),cells).encode('utf-8')
        # Expand ID pickers without shifting any existing records.
        for sheet, part in parts.items():
            raw=changed.get(part,archive.read(part)).decode('utf-8')
            updated=raw.replace(f'$A$6:$A${last}',f'$A$6:$A${last+3}')
            if updated!=raw:changed[part]=updated.encode('utf-8')
        output=BytesIO()
        with ZipFile(output,'w') as target:
            for info in archive.infolist():target.writestr(info,changed.get(info.filename,archive.read(info.filename)))
    candidate=output.getvalue()
    after=openpyxl.load_workbook(BytesIO(candidate),data_only=False)
    cached=openpyxl.load_workbook(BytesIO(candidate),data_only=True)
    for sheet in w:
        allowed_rows=set(new_rows.get(sheet.title,{}))
        for row in sheet:
            for cell in row:
                if cell.row in allowed_rows:continue
                assert after[sheet.title][cell.coordinate].value==cell.value,(sheet.title,cell.coordinate)
        for address,value in summary.get(sheet.title,{}).items():assert cached[sheet.title][address].value==value
    after_ids=[r[0] for r in list(cached['Items'].values)[5:] if r[0]]
    assert after_ids==ids+[i['id'] for i in spec['items']]
    for sheet, rows in new_rows.items():
        for n,(values,caches) in rows.items():
            for c,value in enumerate(values,1):
                assert after[sheet].cell(n,c).value==value,(sheet,n,c)
                assert cached[sheet].cell(n,c).value==caches.get(c,value),(sheet,n,c,'cache')
    for book in [w,v,after,cached]:book.close()
    return candidate,{'sourceHash':digest(snapshot),'resultHash':digest(candidate),'items':[i['id'] for i in spec['items']],'tables':tables,'summaryCaches':summary}


if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    original=SOURCE.read_bytes();(OUT/'before.xlsx').write_bytes(original)
    candidate,receipt=build(original)
    (OUT/'after.xlsx').write_bytes(candidate)
    atomic_write(SOURCE,candidate,receipt['sourceHash'])
    (OUT/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
    print(json.dumps(receipt,indent=2))
