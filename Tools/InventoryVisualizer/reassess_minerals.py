"""One-time mineral value review with editable tier-derived rarity defaults.

Author exact inputs through Artifact Tool, then preserve the master archive with
the existing cell adapter. Never infer crafting readiness from a value tier.
"""
import json
from collections import Counter
from io import BytesIO
from pathlib import Path

import openpyxl
import organize_camp_items as guard
from balance_catalog import NAME_LINKS, SEARCH_FIELDS, merge
from category_edits import verify_move
from tier_edits import INGREDIENTS, tier_number
from workbook_edits import atomic_write, digest, master_items, parse_sheets

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT/'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
SPEC = ROOT/'Docs/Design/Items/MINERAL_VALUE_ASSESSMENT_2026-10-16.json'
AUDIT = ROOT/'Docs/Design/Items/MINERAL_VALUE_AUDIT_2026-10-16.json'
OUT = ROOT/'Logs/MineralValue20261016'
MAPPING = {'T0':'Junk','T1':'Common','T2':'Uncommon','T3':'Rare','T4':'Epic'}


def write_json(path, value):
    path.write_bytes((json.dumps(value, indent=2)+'\n').encode('utf-8'))


def build(snapshot, spec):
    sheets = parse_sheets(snapshot)
    h = sheets['Items'][4]
    records = {r[0]:(n,dict(zip(h,r))) for n,r in master_items(sheets)}
    decisions = {i['id']:i for i in spec['items']}
    assert len(decisions)==len(spec['items'])==50
    assert set(decisions)=={k for k,(_,r) in records.items() if r['Category']=='Mineral resources'}, 'Mineral roster changed; review before applying.'
    changes, updated, reviewed = {}, {}, []

    def change(sheet, number, field, value):
        header = sheets[sheet][4]
        ci = header.index(field)
        row = list(sheets[sheet][number-1])
        if row[ci] != value:
            address = f'{openpyxl.utils.get_column_letter(ci+1)}{number}'
            changes.setdefault(sheet,{})[address] = value
            row[ci] = value
            sheets[sheet][number-1] = tuple(row)

    for key, decision in decisions.items():
        n, old = records[key]
        new = dict(old)
        tier = decision['tier']
        assert tier in MAPPING
        new.update(Tier=tier, Rarity=MAPPING[tier])
        new['Tier basis'] = (f"October 16 mineral value review: {tier}. {decision['reason']} "
            f"Evidence: {', '.join(decision['sources'])}; specific mineral occurrence and grade remain proposals. "
            f"Rarity default: {MAPPING[tier]} from {tier}, independently editable. Value does not certify purity, processing readiness, recipe eligibility, performance or a price.")
        new['Tier sort order'] = tier_number(tier)
        new['Search terms'] = '; '.join(str(new.get(f) or '') for f in SEARCH_FIELDS).lower()
        for field in ['Tier','Rarity','Tier basis','Tier sort order','Search terms']:
            change('Items',n,field,new[field])
        updated[key] = new
        reviewed.append(dict(id=key,name=old['Item name'],previousTier=old['Tier'],tier=tier,
            previousRarity=old['Rarity'],rarity=new['Rarity'],reason=decision['reason'],sources=decision['sources']))

    # Refresh only direct tier projections. Keep names, classification, route
    # topology, quantities, qualification gates and historical readiness intact.
    sources = {'Material tier':'Tier','Current tier':'Tier','Decision basis':'Tier basis',
               'Input tier':'Tier','Output tier':'Tier'}
    for sheet, mappings in NAME_LINKS.items():
        header = sheets[sheet][4]
        for number,row in enumerate(sheets[sheet][5:],6):
            for key_field,field in mappings:
                if field not in {*sources,'Tier','Tier sort order'}:
                    continue
                key = row[header.index(key_field)]
                if key in updated:
                    change(sheet,number,field,updated[key][sources.get(field,field)])

    affected = set()
    rh = sheets['Material routes'][4]
    for number,row in enumerate(sheets['Material routes'][5:],6):
        if row[rh.index('Item ID')] not in updated:
            continue
        ingredient = row[rh.index('Relationship')] in INGREDIENTS
        withheld = row[rh.index('Route status')]=='Withheld'
        numeric = tier_number(row[rh.index('Material tier')]) if ingredient and not withheld else -1
        change('Material routes',number,'Ingredient tier number',numeric)
        change('Material routes',number,'Unresolved ingredient tier',int(ingredient and numeric==-1))
        if ingredient:
            affected.add(row[rh.index('Recipe ID')])
    ch = sheets['Crafting recipes'][4]
    for number,row in enumerate(sheets['Crafting recipes'][5:],6):
        key = row[ch.index('Recipe ID')]
        if key not in affected:
            continue
        values = [r[rh.index('Ingredient tier number')] for r in sheets['Material routes'][5:]
                  if r[rh.index('Recipe ID')]==key and r[rh.index('Relationship')] in INGREDIENTS]
        value = 'Unassigned' if not values or -1 in values else f'T{max(values)}'
        change('Crafting recipes',number,'Listed ingredient tier',value)

    candidate = merge(snapshot,changes,{})
    verify_move(snapshot,candidate,changes,category_contract=False)
    after = parse_sheets(candidate)
    before_items = {r[0]:r for _,r in master_items(parse_sheets(snapshot))}
    after_items = {r[0]:r for _,r in master_items(after)}
    assert before_items.keys()==after_items.keys()
    assert all(before_items[k]==after_items[k] for k in before_items if k not in updated)
    for item in reviewed:
        r = dict(zip(h,after_items[item['id']]))
        assert (r['Tier'],r['Rarity'])==(item['tier'],item['rarity'])
    receipt = dict(sourceHash=digest(snapshot),resultHash=digest(candidate),items=reviewed,
        tiers={t:sum(i['tier']==t for i in reviewed) for t in MAPPING},
        rarities=dict(Counter(i['rarity'] for i in reviewed)),
        changedTiers=sum(i['tier']!=i['previousTier'] for i in reviewed),
        changedRarities=sum(i['rarity']!=i['previousRarity'] for i in reviewed),
        changedCells={s:len(c) for s,c in changes.items()},
        verification='Exact cell allowlist; all other archive parts, cells, styles and formulas preserved; non-mineral master records unchanged.')
    return candidate, changes, receipt


def prepare():
    assert not AUDIT.exists(), 'Already applied; preserve later manual edits.'
    OUT.mkdir(parents=True,exist_ok=True)
    snapshot = SOURCE.read_bytes()
    spec = json.loads(SPEC.read_bytes())
    _, changes, receipt = build(snapshot,spec)
    (OUT/'before.xlsx').write_bytes(snapshot)
    write_json(OUT/'plan.json',dict(changes=changes,receipt=receipt))
    before = {r[0]:(n,r) for n,r in master_items(parse_sheets(snapshot))}
    book = openpyxl.load_workbook(BytesIO(snapshot),data_only=True)
    preview = []
    for tier in MAPPING:
        subset = [i for i in receipt['items'] if i['tier']==tier]
        for item in subset[:2]:
            n,row = before[item['id']]
            preview.append(dict(name=item['name'],tier=item['tier'],rarity=item['rarity'],
                cells=[dict(value=book['Items'][f'{c}{n}'].value,
                    width=book['Items'].column_dimensions[c].width,
                    height=book['Items'].row_dimensions[n].height or 20,
                    font=book['Items'][f'{c}{n}'].font.name or 'Arial',
                    size=book['Items'][f'{c}{n}'].font.sz or 11,
                    bold=bool(book['Items'][f'{c}{n}'].font.b),
                    wrap=bool(book['Items'][f'{c}{n}'].alignment.wrap_text)) for c in ['B','C','P']]))
    book.close()
    write_json(OUT/'preview.json',preview)
    print(json.dumps({k:v for k,v in receipt.items() if k!='items'}))


def apply():
    plan = json.loads((OUT/'plan.json').read_bytes())
    snapshot = (OUT/'before.xlsx').read_bytes()
    assert digest(SOURCE.read_bytes())==digest(snapshot), 'Workbook changed during review; recapture.'
    authored = openpyxl.load_workbook(OUT/'authored.xlsx',data_only=True)
    for sheet,cells in plan['changes'].items():
        for address,value in cells.items():
            assert authored[sheet][address].value==value,(sheet,address,'Authoring mismatch')
    authored.close()
    candidate, changes, receipt = build(snapshot,json.loads(SPEC.read_bytes()))
    assert changes==plan['changes'] and receipt==plan['receipt']
    (OUT/'verified-result.xlsx').write_bytes(candidate)
    guard.release_guard()
    atomic_write(SOURCE,candidate,digest(snapshot))
    write_json(AUDIT,dict(assessment=str(SPEC.relative_to(ROOT)).replace('\\','/'),
        sourceRevision=json.loads(SPEC.read_bytes())['sourceRevision'],liveWorkbook=receipt))
    print('Saved mineral review and verified exact preservation.')


if __name__=='__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('step',choices=['prepare','apply','run'])
    step = parser.parse_args().step
    if step=='run':
        guard.OUT=OUT
        guard.prepare=prepare
        guard.apply=apply
        guard.run()
    else:
        globals()[step]()
