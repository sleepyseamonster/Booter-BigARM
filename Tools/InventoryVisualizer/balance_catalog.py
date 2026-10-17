"""Preserving value-tier and identity audit, excluding both salvage categories.

Preparation reads the reviewed assessment/refinements from the local operation
directory. Apply can rebase over concurrent changes to untouched cells, but
rejects any conflict in a task-owned cell or any unreviewed eligible item.
"""
import argparse
import json
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from category_edits import CELL, parts_for, patch_cells, verify_move
from tier_edits import INGREDIENTS, tier_number
from workbook_edits import atomic_write, digest, master_items, parse_sheets

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT/'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
OUT = ROOT/'Logs/CatalogBalance20261016'
AUDIT = ROOT/'Docs/Design/Items/CATALOG_VALUE_BALANCE_2026-10-16.json'
EXCLUDED = {'Salvage', 'Ancient salvage'}
SEARCH_FIELDS = ['Item ID','Item name','Category','Subtype','Item type','Processing stage',
                 'Origin family','Primary use','Material family','Physical form','Design status',
                 'Search aliases','Tier','Finished-item input tier','Crafting role']
NAME_LINKS = {
    'Item uses':[('Item ID','Item name')],
    'Material routes':[('Item ID','Item name'),('Item ID','Item type'),('Item ID','Material tier')],
    'Tier review':[('Item ID','Item name'),('Item ID','Current tier'),('Item ID','Tier sort order'),('Item ID','Decision basis')],
    'Craft audit':[('Item ID','Item name'),('Item ID','Item type'),('Item ID','Tier')],
    'Crafting browser':[('Input item ID','Input name'),('Output item ID','Output name'),
                        ('Input item ID','Input tier'),('Output item ID','Output tier'),
                        ('Input item ID','Input type'),('Output item ID','Output type')],
    'Item relationships':[('Parent item ID','Parent item'),('Related item ID','Related item')],
}


def merge(snapshot, changes, formulas):
    stream = BytesIO()
    with ZipFile(BytesIO(snapshot)) as archive, ZipFile(stream,'w') as result:
        parts = parts_for(archive)
        replacement = {}
        for sheet, edits in changes.items():
            raw = patch_cells(archive.read(parts[sheet]).decode(), edits)
            selected = formulas.get(sheet,{})
            def formula_cell(match):
                address = match.group(1)
                cell = match.group()
                if address not in selected:
                    return cell
                expression = selected[address]
                tag = re.search(r'<(?:\w+:)?f\b[^>]*(?:/>|>.*?</(?:\w+:)?f>)',cell,re.S)
                assert tag, (sheet,address,'expected formula')
                from xml.sax.saxutils import escape
                prefix = re.search(r'<((?:\w+:)?)c\b',cell).group(1)
                return cell.replace(tag.group(), f'<{prefix}f>{escape(expression)}</{prefix}f>',1)
            raw = re.sub(CELL,formula_cell,raw,flags=re.S)
            replacement[parts[sheet]] = raw.encode()
        for info in archive.infolist():
            result.writestr(info,replacement.get(info.filename,archive.read(info.filename)))
    return stream.getvalue()


def verify(snapshot, candidate, changes, formulas):
    # Existing verifier proves exact preservation of all cells/styles/XML parts.
    # Give it a source with only the explicitly authorized formula replacements.
    original_sheets = parse_sheets(snapshot)
    normalized = merge(snapshot, {s:{a:original_sheets[s][int(re.search(r'\d+',a).group())-1][openpyxl.utils.cell.column_index_from_string(re.match(r'[A-Z]+',a).group())-1]
                                        for a in f} for s,f in formulas.items()}, formulas) if formulas else snapshot
    verify_move(normalized,candidate,changes,category_contract=False)
    with ZipFile(BytesIO(snapshot)) as before, ZipFile(BytesIO(candidate)) as after:
        assert before.read('xl/workbook.xml') == after.read('xl/workbook.xml')
    old = original_sheets
    new = parse_sheets(candidate)
    old_items = {r[0]:r for _,r in master_items(old)}
    new_items = {r[0]:r for _,r in master_items(new)}
    assert list(old_items)==list(new_items)
    assert Counter(r[3] for r in old_items.values())==Counter(r[3] for r in new_items.values())
    excluded = {k for k,r in old_items.items() if r[3] in EXCLUDED}
    assert all(old_items[k]==new_items[k] for k in excluded), 'Protected salvage record changed'
    # No row that directly references a protected item may be modified, including
    # mixed input/output browser rows and relationship records.
    for sheet, edits in changes.items():
        for address in edits:
            n = int(re.search(r'\d+',address).group())
            assert not any(isinstance(v,str) and v in excluded for v in old[sheet][n-1]), (sheet,address,'salvage reference')


def prepare():
    assert not AUDIT.exists(), 'Already applied; inspect the receipt rather than replaying.'
    original = SOURCE.read_bytes()
    sheets = parse_sheets(original)
    header = sheets['Items'][4]
    byid = {r[0]:(n,dict(zip(header,r))) for n,r in master_items(sheets)}
    assert len(byid)==len(master_items(sheets))
    excluded = {k for k,(_,r) in byid.items() if r['Category'] in EXCLUDED}
    assessment = {i['id']:i for i in json.loads((OUT/'assessment.json').read_text())}
    assert set(assessment)==set(byid)-excluded, 'New, removed or reclassified eligible items need review.'
    refinements = {i['id']:i for i in json.loads((OUT/'refinements.json').read_text())}
    assert set(refinements)<=set(assessment)
    formula_book = openpyxl.load_workbook(BytesIO(original),data_only=False)
    changes, formulas, reviewed = {}, {}, []

    def change(sheet,n,field,value,formula=None):
        h = sheets[sheet][4]
        ci = h.index(field)
        address = f'{openpyxl.utils.get_column_letter(ci+1)}{n}'
        row = list(sheets[sheet][n-1])
        if row[ci]!=value or formula is not None:
            changes.setdefault(sheet,{})[address] = value
            row[ci] = value
            sheets[sheet][n-1] = tuple(row)
        if formula is not None and formula_book[sheet][address].data_type=='f':
            formulas.setdefault(sheet,{})[address] = formula

    updated = {}
    for key, item in assessment.items():
        n, old = byid[key]
        new = dict(old)
        refinement = refinements.get(key,{})
        new.update(refinement.get('fields',{}))
        new['Tier'] = item['tier']
        if old['Item name']!=new['Item name']:
            new['Search aliases'] = '; '.join(dict.fromkeys(v for v in [old.get('Search aliases'),old['Item name']] if v))
        new['Tier basis'] = f"October 16 catalog value audit: {new['Tier']}. {item['reason']} Value is separate from crafting readiness, rarity, condition and demonstrated performance."
        if old.get('Design status')=='Withheld assumption':
            new['Tier basis'] += ' This unresolved source remains withheld; its numeric rank establishes no obtainable resource or saleable yield.'
        new['Tier sort order'] = tier_number(new['Tier'])
        new['Search terms'] = '; '.join(str(new.get(k) or '') for k in SEARCH_FIELDS).lower()
        new['Crafting coverage'] = ('Route listed; review status/completeness' if (new.get('Listed output routes') or 0)>0 else 'No production route listed')
        for field in ['Tier','Tier basis','Tier sort order','Search aliases','Search terms',*refinement.get('fields',{})]:
            change('Items',n,field,new.get(field))
        change('Items',n,'Crafting coverage',new['Crafting coverage'],f'IF(AJ{n}>0,"Route listed; review status/completeness","No production route listed")')
        updated[key] = new
        reviewed.append(dict(id=key,category=old['Category'],previousName=old['Item name'],name=new['Item name'],
                             previousTier=old['Tier'],tier=new['Tier'],valueReason=item['reason'],
                             refinements={f:dict(previous=old.get(f),value=v) for f,v in refinement.get('fields',{}).items()},
                             refinementReason=refinement.get('reason')))
    assert len({r['Item name'].casefold() for r in updated.values()})==len(updated), 'Duplicate revised names'
    assert not ({r['Item name'].casefold() for r in updated.values()} & {byid[k][1]['Item name'].casefold() for k in excluded})
    for sheet,mappings in NAME_LINKS.items():
        h = sheets[sheet][4]
        for n,row in enumerate(sheets[sheet][5:],6):
            if any(isinstance(v,str) and v in excluded for v in row):
                continue
            for key_field,field in mappings:
                key = row[h.index(key_field)]
                if key not in updated:
                    continue
                source = {'Material tier':'Tier','Current tier':'Tier','Decision basis':'Tier basis',
                          'Input name':'Item name','Output name':'Item name','Parent item':'Item name','Related item':'Item name',
                          'Input tier':'Tier','Output tier':'Tier','Input type':'Item type','Output type':'Item type'}.get(field,field)
                change(sheet,n,field,updated[key][source])
    # Tier fields in the existing linked schema now describe value, never recipe
    # legality. Preserve route topology, quantities, source/withheld gates and all
    # rows involving protected salvage. Display this distinction after Excel recalcs.
    h = sheets['Crafting browser'][4]
    for n,row in enumerate(sheets['Crafting browser'][5:],6):
        if any(isinstance(v,str) and v in excluded for v in row):
            continue
        if row[h.index('Input item ID')] in updated or row[h.index('Output item ID')] in updated:
            value = ('Withheld premise' if row[h.index('Route status')]=='Withheld' else
                     'Reusable prerequisite; value tier independent' if row[h.index('Relationship')]=='Tool candidate' else
                     'Value tier; crafting readiness separate')
            formula = f'IF(L{n}="Withheld","Withheld premise",IF(E{n}="Tool candidate","Reusable prerequisite; value tier independent","Value tier; crafting readiness separate"))'
            change('Crafting browser',n,'Tier interpretation',value,formula)
    # Existing schema names are retained for compatibility. Linked numeric tier
    # helpers project value ranks and never override the route's qualification.
    rh = sheets['Material routes'][4]
    excluded_recipes = {r[rh.index('Recipe ID')] for r in sheets['Material routes'][5:]
                        if r[rh.index('Item ID')] in excluded}
    affected_recipes = set()
    for n,row in enumerate(sheets['Material routes'][5:],6):
        if row[rh.index('Item ID')] not in updated:
            continue
        ingredient = row[rh.index('Relationship')] in INGREDIENTS
        withheld = row[rh.index('Route status')]=='Withheld'
        numeric = tier_number(row[rh.index('Material tier')]) if ingredient and not withheld else -1
        original_formula = formula_book['Material routes'][f'P{n}'].value
        withheld_formula = f'IF(K{n}="Withheld",-1,{original_formula[1:]})' if withheld and isinstance(original_formula,str) and original_formula.startswith('=') else None
        change('Material routes',n,'Ingredient tier number',numeric,withheld_formula)
        change('Material routes',n,'Unresolved ingredient tier',int(ingredient and numeric==-1))
        if ingredient:
            affected_recipes.add(row[rh.index('Recipe ID')])
    ch = sheets['Crafting recipes'][4]
    for n,row in enumerate(sheets['Crafting recipes'][5:],6):
        key = row[ch.index('Recipe ID')]
        if key not in affected_recipes or key in excluded_recipes:
            continue
        ingredients = [r for r in sheets['Material routes'][5:] if r[rh.index('Recipe ID')]==key and r[rh.index('Relationship')] in INGREDIENTS]
        values = [r[rh.index('Ingredient tier number')] for r in ingredients]
        value = 'Unassigned' if not values or -1 in values else f'T{max(values)}'
        change('Crafting recipes',n,'Listed ingredient tier',value)
    names = {i['previousName']:i['name'] for i in reviewed if i['previousName']!=i['name']}
    pattern = re.compile(r'(?<![\w])(?:'+'|'.join(re.escape(k) for k in sorted(names,key=len,reverse=True))+r')(?![\w])')
    # Refresh cached name lists only in selected display views; keep imported
    # Variants and authored Unity asset descriptions as original source facts.
    for sheet,fields in {'Tier review':['Candidate input materials'], 'Craft audit':['Material coverage','Facilities / reusable tools']}.items():
        h = sheets[sheet][4]
        for n,row in enumerate(sheets[sheet][5:],6):
            if any(isinstance(v,str) and v in excluded for v in row):
                continue
            for field in fields:
                value = row[h.index(field)]
                if isinstance(value,str):
                    change(sheet,n,field,pattern.sub(lambda m:names[m.group()],value))
    for n,row in enumerate(sheets['Guide'][5:],6):
        replacements = {
            'Tier versus rarity and quality':'Current catalog tiers outside Salvage and Ancient salvage rank estimated practical worth: T0 common, T1 serviceable, T2 specialist, T3 high value and T4 exceptional. Replacement labor, utility, scarcity and condition inform the ranking. They do not establish recipe eligibility, manufacturing stages, numeric prices or power. Both salvage sections are excluded from this audit.',
            'Tier fields':'Items.C is the current value rank T0-T4; AG records the reason. Historical Previous tier and finished-input fields retain earlier manufacturing-planning information. Every item outside Salvage and Ancient salvage is assigned a numeric value tier. Withheld sources remain withheld even when a provisional value tier is present.',
            'Crafting tier framework':'The October 15 refinement-stage framework is retained as historical manufacturing-planning guidance. October 16 item value tiers supersede Items.C outside the two excluded salvage sections. A high-value raw item can require preparation; a low-value finished fitting can be immediately usable after qualification. Price bands cannot validate a material-tier progression.',
            'Recipe scope and tier gates':'Material routes and Crafting recipes retain candidate topology, quantities and source/withheld gates. Their existing linked tier columns project value ranks for items outside both salvage sections; these legacy field labels do not validate a refinement-stage progression. Crafting-browser rows explicitly separate value from readiness. No ingredient/output permission follows from a numeric value rank; manufacturing stages need their own separate review.',
        }
        if row[0] in replacements:
            change('Guide',n,'Guidance',replacements[row[0]])
    for n,row in enumerate(sheets['Taxonomy'][5:],6):
        if row[0]=='Material tier':
            text = ('Legacy schema label for the item tier selector. Outside Salvage and Ancient salvage, T0-T4 are relative practical-worth bands from common to exceptional, independent of readiness, quality and demonstrated function. Not applicable and Unassigned remain historical selector options. Withheld source status remains authoritative.')
            change('Taxonomy',n,'Definition / rule',text)
    formula_book.close()
    counts = Counter(i['tier'] for i in reviewed)
    assert all(counts[f'T{i}']>counts[f'T{i+1}'] for i in range(4))
    plan = dict(sourceHash=digest(original),changes=changes,formulas=formulas,items=reviewed,
                counts=dict(sorted(counts.items())),excludedCategories=sorted(EXCLUDED))
    (OUT/'before.xlsx').write_bytes(original)
    (OUT/'plan.json').write_text(json.dumps(plan,indent=2)+'\n')
    # Ten representative value rows, including refinements, for scoped layout QA.
    samples=[]
    for t in range(5):
        subset=[i for i in reviewed if i['tier']==f'T{t}']
        samples.extend(subset[:1]+[next((i for i in subset if i['refinements']),subset[-1])])
    book=openpyxl.load_workbook(BytesIO(original),data_only=True)
    preview=[]
    for item in samples:
        n=byid[item['id']][0]
        preview.append(dict(cells=[dict(value=book['Items'][f'{col}{n}'].value,
            width=book['Items'].column_dimensions[col].width,height=book['Items'].row_dimensions[n].height or 20,
            font=book['Items'][f'{col}{n}'].font.name or 'Arial',size=book['Items'][f'{col}{n}'].font.sz or 11,
            bold=bool(book['Items'][f'{col}{n}'].font.b),wrap=bool(book['Items'][f'{col}{n}'].alignment.wrap_text)) for col in ['B','C']],
            name=item['name'],tier=item['tier']))
    book.close()
    (OUT/'preview.json').write_text(json.dumps(preview,indent=2))
    print('Prepared',len(reviewed),'items',dict(counts),'refined',sum(bool(i['refinements']) for i in reviewed))


def apply():
    assert not AUDIT.exists()
    plan=json.loads((OUT/'plan.json').read_text())
    original=(OUT/'before.xlsx').read_bytes()
    latest=SOURCE.read_bytes()
    # Permit live salvage/review edits, rejecting conflicting task-owned edits.
    baseline=parse_sheets(original)
    current=parse_sheets(latest)
    eligible={r[0] for _,r in master_items(current) if r[3] not in EXCLUDED}
    assert eligible=={i['id'] for i in plan['items']}, 'Eligible roster changed; review the new items before applying.'
    for sheet,edits in plan['changes'].items():
        for address in edits:
            n=int(re.search(r'\d+',address).group())-1
            ci=openpyxl.utils.cell.column_index_from_string(re.match(r'[A-Z]+',address).group())-1
            assert current[sheet][n][ci]==baseline[sheet][n][ci], (sheet,address,'Concurrent task-owned change; recapture first')
    authored=openpyxl.load_workbook(OUT/'authored.xlsx',data_only=True)
    for sheet,edits in plan['changes'].items():
        for address,value in edits.items():
            assert authored[sheet][address].value==value, (sheet,address)
    authored.close()
    candidate=merge(latest,plan['changes'],plan['formulas'])
    verify(latest,candidate,plan['changes'],plan['formulas'])
    audit=dict(date='2026-10-16',authority='Author requested a coherent practical-worth audit, names/identity refinements where useful, and descending tier counts. Author then explicitly excluded salvage.',
        scope='Catalog-only local workbook. Both Salvage and Ancient salvage and all directly referencing rows are unchanged. No Google Sheets synchronization, Unity asset changes or external publication.',
        method='Each eligible item was assessed by practical utility, source/condition uncertainty, replaceability, prepared-material scarcity, fabrication effort and plausible specialist demand. A second pass compared stock, fitted parts and whole assemblies. T0 is largest and T4 smallest; no equal per-category quota was imposed.',
        limits='Value bands are planning judgments, not sale prices, official currency exchange rates, power ratings or manufacturing-stage permissions. Withheld biology stays withheld. Historical route topology, quantities and manufacturing-stage records are retained. Old names remain searchable aliases; Variants retain authored Unity source facts.',
        worldBoundary='World identity, streaming/unload/reload, generated-object identity, authored placement and persisted runtime deltas are not applicable: this task creates no runtime objects and changes no save or world data.',
        counts=plan['counts'],reviewed=len(plan['items']),refined=sum(bool(i['refinements']) for i in plan['items']),
        renamed=sum(i['name']!=i['previousName'] for i in plan['items']),items=plan['items'],
        sourceHashAtPreparation=plan['sourceHash'],sourceHashBefore=digest(latest),sourceHashAfter=digest(candidate),
        verification=dict(idsAndOrderPreserved=True,categoriesPreserved=True,salvageMasterAndDirectReferencesUnchanged=True,
            originalFormulasAndStyles='Preserved except explicit value-semantics formulas in Crafting coverage and Tier interpretation and the withheld-source guard in route numeric helpers.',
            changedCells={s:len(v) for s,v in plan['changes'].items()},formulaReplacements={s:len(v) for s,v in plan['formulas'].items()}))
    (OUT/'pre-apply.xlsx').write_bytes(latest)
    (OUT/'candidate.xlsx').write_bytes(candidate)
    atomic_write(SOURCE,candidate,digest(latest))
    atomic_write(AUDIT,(json.dumps(audit,indent=2)+'\n').encode())
    print('Applied',audit['reviewed'],'items:',audit['counts'],'refined',audit['refined'],'renamed',audit['renamed'])


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('action',choices=['prepare','apply'])
    prepare() if parser.parse_args().action=='prepare' else apply()
