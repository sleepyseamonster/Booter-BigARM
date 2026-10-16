"""Create five five-piece armor families; never create a counted whole-set item."""
import argparse
import json
import subprocess
import ctypes
from ctypes import wintypes
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from add_ancient_salvage import ROOT, SOURCE, build
from category_edits import parts_for, patch_cells, verify_move
from workbook_edits import atomic_write, digest, parse_sheets, build_deletion

OUT=ROOT/'Logs/FiveArmorFamilies20261015'
DATA=ROOT/'Docs/Design/Items/FIVE_ARMOR_FAMILIES_2026-10-15.json'
MANIFEST=ROOT/'Docs/Design/Items/EQUIPMENT_CLASSIFICATION.json'
AGGREGATE='proposal.gear.metal_armor_set'
_write_guard=None


def release_guard():
    global _write_guard
    if _write_guard is not None:
        kernel,handle=_write_guard
        kernel.CloseHandle(handle)
        _write_guard=None


def run_guarded():
    """Allow readers while briefly excluding independent workbook writers."""
    global _write_guard
    kernel=ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.CreateFileW.argtypes=[wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,wintypes.LPVOID,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
    kernel.CreateFileW.restype=wintypes.HANDLE
    kernel.CloseHandle.argtypes=[wintypes.HANDLE]
    kernel.CloseHandle.restype=wintypes.BOOL
    handle=kernel.CreateFileW(str(SOURCE),0x80000000,1,None,3,0x80,None)
    if handle==ctypes.c_void_p(-1).value:raise ctypes.WinError(ctypes.get_last_error())
    _write_guard=(kernel,handle)
    try:
        prepare()
        node=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
        subprocess.run([str(node),str(OUT/'stage.mjs'),str(OUT)],check=True)
        apply()
    finally:release_guard()
REUSE={'chest':'proposal.gear.chest_armor','legs':'proposal.gear.leg_armor','helmet':'proposal.gear.helmet'}
FAMILIES=[
    ('iron','Iron','Metal','T2','T1','Broad worked-iron plates with plain rolled edges and exposed metal fasteners.'),
    ('steel','Steel','Metal','T3','T2','Closer-fitting steel plates with overlapping articulation and narrow reinforced edges.'),
    ('chitin','Chitin','Shell / chitin','T2','T1','Layered curved shell plates with visible growth texture and small metal or tendon joining points.'),
    ('ceramic','Ceramic-faced','Ceramic and metal','T3','T2','Small replaceable ceramic faces carried on an articulated metal support, with visible gaps between tiles.'),
    ('plasteel','Ancient plasteel','Ancient composite','T0','T0','Recovered pale composite segments with unfamiliar recessed fittings, wear and mismatched surviving surfaces.'),
]
PIECES=[('chest','chestplate','Torso and upper back; a chest piece does not include arm or leg pieces.'),
        ('legs','leg guards','Paired thigh, knee and shin guards; boots remain a separate item.'),
        ('boots','armored boots','A matched pair of armored foot coverings, separate from shin guards.'),
        ('arms','arm guards','Paired upper-arm, elbow and forearm guards; no extra glove item or hand-protection claim.'),
        ('helmet','helmet','Head covering with sight, hearing and canister clearance requiring design.')]


def specification():
    defaults={'Tier':'T2','Category':'Armor','Subtype':'Armor family piece','Implementation scope':'New proposal',
        'Equipment slot':'Unassigned; no new slot approved','Carrier metadata':'Booter-worn armor concept; fitting, packing and spare-cargo rules unresolved.',
        'Stack limit':None,'Mass (game units)':None,'Acquisition':'Proposed contemporary fabrication or specialist trade; construction recipe unresolved.',
        'Crafting / processing':'Separate fitted piece using selected qualified stock; no whole-set recipe, quantities, yield or production time authored.',
        'Supported behavior':'Not implemented','Limitations':'Fit, coverage, articulation, mass, durability and protection remain proposals. No wood, ordinary cloth, free-water quenching, set bonuses or new equipment slots.',
        'Rarity':None,'Item type':'Armor','Processing stage':'Finished item','Origin family':'Manufactured / unresolved',
        'Primary use':'Protection','Material family':'Metal','Physical form':'Assembly','Design status':'Proposal',
        'Classification basis':'One independently listed armor piece. Family membership is a design grouping, never another inventory item or a set bonus.',
        'Tier basis':'Proposed finished-piece planning tier; recipe and properties unresolved.','Finished-item input tier':'T1','Crafting role':'Finished item'}
    items=[];groups=[]
    for family,label,material,tier,input_tier,appearance in FAMILIES:
        ids=[]
        for piece,noun,coverage in PIECES:
            key=REUSE[piece] if family=='iron' and piece in REUSE else f'proposal.armor.{family}.{piece}'
            name=f'{label} {noun}'
            fields={'Tier':tier,'Finished-item input tier':input_tier,'Material family':material,
                'Subtype':f'{label} armor / {piece}',
                'Search aliases':f'{family} armor set; {family} armor family; {label.lower()} armor; {piece}; separate armor piece',
                'Tier basis':f'Proposed {tier} finished piece using {input_tier} inputs; material grade, fabrication and protection are unresolved.',
                'Open design decisions':f'{coverage} Visual direction: {appearance} Joining, fit, wearability and performance require design.'}
            purpose=f'{label} armor family: individually listed {noun}. {coverage} {appearance} Construction and protection are proposals; the five family pieces remain separate items.'
            if family=='plasteel':
                fields.update({'Origin family':'Ancient technology','Processing stage':'Recovered part','Rarity':'Extremely rare',
                    'Acquisition':'Recover a surviving ancient armor piece from a proposed ruin or machine-age cache; sites and availability unresolved.',
                    'Crafting / processing':'Inspect and qualify a surviving original for fit and integrity. No contemporary plasteel manufacturing, internal repair or complete-set acquisition route established.',
                    'Tier basis':'T0 as found, independent of ancient complexity and extreme rarity. A compatible original would need a separate T1 qualification record before reuse.',
                    'Limitations':'Extremely rare ancient recovery. Fit, integrity, joining and protection are unknown. No contemporary plasteel manufacture, guaranteed usable condition, set bonus or working protective effect.'})
                purpose=f'Extremely rare recovered ancient plasteel {noun}, listed as one separate armor item. {coverage} {appearance} Presence does not establish contemporary manufacture or a usable fitted piece.'
            items.append(dict(id=key,name=name,purpose=purpose,icon=f'{name}: {appearance} Show only this piece, not a complete armor suit.',material=material,use='Protection',fields=fields))
            ids.append(key)
        groups.append(dict(family=family,label=label,itemIds=ids,itemCount=5))
    return dict(source='Author requested five armor types with separate chest, legs, boots, arms and helmet entries, October 15, 2026. Plasteel, iron and chitin were requested examples; steel and ceramic-faced armor complete this first brainstorm. Docs/WORLD_BASIS.md controls material and ancient-recovery boundaries.',
        scope='25 individual piece records across five design families. No whole-set inventory record, automatic full-set grant, set bonus, gameplay equipment slots, world placement, streaming or saved runtime deltas.',
        retiredAggregate=AGGREGATE,reusedIds=list(REUSE.values()),defaults=defaults,groups=groups,items=items)


def prepare():
    OUT.mkdir(parents=True,exist_ok=True)
    snapshot=SOURCE.read_bytes()
    spec=specification()
    rows=parse_sheets(snapshot)['Items'][5:]
    assert {r[0] for r in rows if r[0]} >= {AGGREGATE,*REUSE.values()}
    if DATA.exists():assert json.loads(DATA.read_text(encoding='utf-8'))==spec,'Specification changed; review before preparing'
    else:DATA.write_text(json.dumps(spec,indent=2)+'\n',encoding='utf-8',newline='\n')
    (OUT/'before.xlsx').write_bytes(snapshot)
    (OUT/'manifest-before.json').write_bytes(MANIFEST.read_bytes())
    (OUT/'plan.json').write_text(json.dumps(dict(sourceHash=digest(snapshot),spec=spec),indent=2),encoding='utf-8')
    print('Prepared 25 pieces: reuse 3, append 22, retire 1 whole-set record.')


def apply():
    plan=json.loads((OUT/'plan.json').read_text(encoding='utf-8'));spec=plan['spec']
    assert json.loads(DATA.read_text(encoding='utf-8'))==spec and digest(SOURCE.read_bytes())==plan['sourceHash']
    authored=openpyxl.load_workbook(OUT/'authored.xlsx',read_only=True,data_only=True)
    rows=list(authored['Catalog inputs'].values)[1:]
    assert len(rows)==25
    for row,i in zip(rows,spec['items']):
        assert list(row[:6])==[i[k] for k in ['id','name','purpose','icon','material','use']] and json.loads(row[6])==i['fields']
    authored.close()
    snapshot=(OUT/'before.xlsx').read_bytes()
    candidate,deletion=build_deletion(snapshot,AGGREGATE)
    sheets=parse_sheets(candidate);h=sheets['Items'][4]
    edits={};reused={i['id']:i for i in spec['items'] if i['id'] in REUSE.values()}
    for key,item in reused.items():
        n,old=next((n,r) for n,r in enumerate(sheets['Items'],1) if r and r[0]==key)
        new=dict(spec['defaults']);new.update(item['fields']);new.update({'Item ID':key,'Item name':item['name'],'Purpose':item['purpose'],'Sources':spec['source']})
        new['Search terms']='; '.join(str(new.get(k) or '') for k in ['Item ID','Item name','Category','Subtype','Item type','Processing stage','Origin family','Primary use','Material family','Physical form','Design status','Search aliases','Tier','Finished-item input tier','Crafting role']).lower()
        new['Tier sort order']=2
        for field,value in new.items():
            if field in h:edits.setdefault('Items',{})[f'{openpyxl.utils.get_column_letter(h.index(field)+1)}{n}']=value if value is not None else ''
        for sheet,rs in sheets.items():
            if sheet=='Items' or len(rs)<5:continue
            header=rs[4]
            for number,row in enumerate(rs[5:],6):
                if key not in row:continue
                for ci,field in enumerate(header):
                    if field in ['Item name','Tier','Item type','Processing stage','Origin family','Material family','Primary use','Tier basis','Finished-item input tier','Tier sort order'] and field in new and ci<len(row):
                        edits.setdefault(sheet,{})[f'{openpyxl.utils.get_column_letter(ci+1)}{number}']=new[field]
    with ZipFile(BytesIO(candidate)) as archive:
        parts=parts_for(archive);changed={parts[s]:patch_cells(archive.read(parts[s]).decode(),v).encode() for s,v in edits.items()};out=BytesIO()
        with ZipFile(out,'w') as z:
            for info in archive.infolist():z.writestr(info,changed.get(info.filename,archive.read(info.filename)))
    revised=out.getvalue()
    (OUT/'metadata-before.xlsx').write_bytes(candidate)
    (OUT/'metadata-after.xlsx').write_bytes(revised)
    verify_move(candidate,revised,edits)
    append_spec=dict(spec);append_spec['items']=[i for i in spec['items'] if i['id'] not in reused]
    final,receipt=build(revised,append_spec)
    result=parse_sheets(final)['Items'][5:];byid={r[0]:r for r in result if r[0]}
    assert AGGREGATE not in byid
    allids=[i for g in spec['groups'] for i in g['itemIds']]
    assert len(allids)==len(set(allids))==25 and all(key in byid for key in allids)
    assert all(byid[i][3]=='Armor' and byid[i][23]=='Armor' for i in allids)
    old_manifest=(OUT/'manifest-before.json').read_bytes();assert MANIFEST.read_bytes()==old_manifest
    (OUT/'after.xlsx').write_bytes(final)
    release_guard()
    atomic_write(SOURCE,final,plan['sourceHash'])
    manifest=json.loads(old_manifest)
    manifest['items']=[i for i in manifest['items'] if i['id']!=AGGREGATE]
    for item in manifest['items']:
        if item['id'] in reused:item.update(name=reused[item['id']]['name'],reason='Generic armor piece refined into the iron family; ID retained.')
    manifest['counts']=dict(Counter(i['category'] for i in manifest['items']))
    manifest['sourceHashAfter']=digest(final)
    manifest.setdefault('revisions',[]).append(dict(date='2026-10-15',reason='Replace counted whole-set concept with five individually listed armor families.',retiredIds=[AGGREGATE],reusedIds=list(reused),sourceHashBefore=plan['sourceHash'],sourceHashAfter=digest(final)))
    atomic_write(MANIFEST,(json.dumps(manifest,indent=2)+'\n').encode(),digest(old_manifest))
    receipt.update(sourceHash=plan['sourceHash'],resultHash=digest(final),retiredAggregate=AGGREGATE,reusedIds=list(reused),pieceCount=25,deletion=deletion)
    (OUT/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
    print('Verified: five families, 25 separate pieces, no whole-set item.')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('action',choices=['prepare','apply','run'])
    action=p.parse_args().action
    run_guarded() if action=='run' else prepare() if action=='prepare' else apply()
