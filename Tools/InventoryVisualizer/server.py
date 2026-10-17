"""Loopback-only workbook visualizer with explicit item review actions."""
import argparse
import hashlib
import json
import secrets
from urllib.parse import urlsplit
from io import BytesIO
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import workbook_edits
import category_edits
import tier_edits
import rarity_edits

ROOT = Path(__file__).resolve().parents[2]
WEB = Path(__file__).resolve().parent / "dist"
WORKBOOK = ROOT / "Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx"
REFERENCE = ROOT / "Docs/Agents/Sable/work/2026-10-15-material-dioramas/refined_BC/yard_E_sunset_close.jpg"
REVIEW = ROOT / "Docs/Design/Items/ITEM_REVIEW_STATE.json"
BACKUPS = ROOT / "Logs/InventoryVisualizer/Backups"
TOKEN = secrets.token_urlsafe(32)


def review_metadata():
    state = workbook_edits.review_state(REVIEW)
    return {"protectedIds": state["protectedIds"],
            "reviewIds": state["reviewIds"],
            "reviewRevision": hashlib.sha256(json.dumps(state, sort_keys=True).encode()).hexdigest(),
            "mutationToken": TOKEN}


def catalog():
    snapshot = WORKBOOK.read_bytes()
    book = openpyxl.load_workbook(BytesIO(snapshot), read_only=True, data_only=True)
    try:
        rows = iter(book["Items"].values)
        headers = next(row for row in rows if row and row[0] == "Item ID")
        items = [dict((key, value) for key, value in zip(headers, row) if key)
                 for row in rows if row and row[0]]
        ids = [item["Item ID"] for item in items]
        if len(ids) != len(set(ids)):
            raise ValueError("Duplicate item IDs in master workbook")
        categories = (category_edits.registry(list(book['Categories'].values)) if 'Categories' in book.sheetnames
                      else sorted({item['Category'] for item in items if item.get('Category')}))
        tiers = tier_edits.tier_choices({'Taxonomy': list(book['Taxonomy'].values)}) if 'Taxonomy' in book.sheetnames else []
        return {"items": items, "categories": categories, "categoryEditing": 'Categories' in book.sheetnames,
                "tierChoices": tiers, "tierEditing": bool(tiers),
                "rarityChoices": rarity_edits.RARITIES, "rarityEditing": 'Rarity' in headers,
                "source": str(WORKBOOK.relative_to(ROOT)),
                "sha256": hashlib.sha256(snapshot).hexdigest(), **review_metadata()}
    finally:
        book.close()


class Handler(SimpleHTTPRequestHandler):
    def json_response(self, status, data):
        payload = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def do_POST(self):
        # Explicit same-origin JSON actions only; other sites cannot issue mutations.
        host = self.headers.get("Host", "")
        origin = self.headers.get("Origin", "")
        if (host not in {f"127.0.0.1:{self.server.server_port}", f"localhost:{self.server.server_port}"}
                or origin != "http://" + host
                or not secrets.compare_digest(self.headers.get("X-Inventory-Token", ""), TOKEN)
                or self.headers.get("Content-Type", "").split(";")[0] != "application/json"):
            self.json_response(403, {"error": "Reload this local page before changing items."})
            return
        path = urlsplit(self.path).path
        if path not in {"/api/items/protect", "/api/items/review", "/api/items/delete", "/api/items/category", "/api/items/tier", "/api/items/rarity"}:
            self.json_response(404, {"error": "Unknown action"})
            return
        try:
            length = int(self.headers.get("Content-Length", 0))
            if not 0 < length <= 4096:
                raise workbook_edits.EditError("Invalid request size.", 400)
            body = json.loads(self.rfile.read(length))
            if not isinstance(body, dict):
                raise workbook_edits.EditError("A JSON object is required.", 400)
            if not isinstance(body.get("itemId"), str) or not isinstance(body.get("sourceHash"), str):
                raise workbook_edits.EditError("Item ID and spreadsheet revision are required.", 400)
            if path == '/api/items/rarity':
                if not all(isinstance(body.get(key), str) for key in ['fromRarity', 'toRarity']):
                    raise workbook_edits.EditError('Source and destination rarities are required.', 400)
                result = rarity_edits.update(WORKBOOK, BACKUPS, body)
            elif path == '/api/items/tier':
                if not all(isinstance(body.get(key),str) for key in ['fromTier','toTier']):
                    raise workbook_edits.EditError('Source and destination tiers are required.',400)
                result = tier_edits.update(WORKBOOK, BACKUPS, body)
            elif path == '/api/items/category':
                if not all(isinstance(body.get(key), str) for key in ['toCategory', 'fromCategory']):
                    raise workbook_edits.EditError("Source and destination categories are required.", 400)
                result = category_edits.move(WORKBOOK, BACKUPS, body)
            else:
                result = workbook_edits.mutate(WORKBOOK, REVIEW, BACKUPS, path.rsplit("/", 1)[1],
                                               body["itemId"], body["sourceHash"], body.get("protected"), body.get("reviewed"))
            self.json_response(200, {**result, **review_metadata()})
        except workbook_edits.EditError as error:
            self.json_response(error.status, {"error": str(error)})
        except (PermissionError, OSError):
            self.json_response(423, {"error": "Spreadsheet is locked or unavailable. Close it in Excel, reload and try again."})
        except Exception:
            self.json_response(400, {"error": "The action could not be completed. Reload the spreadsheet and try again."})

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        super().end_headers()

    def do_GET(self):
        path = self.path.split("?", 1)[0]
        if path.startswith('/api/rarity-operations/'):
            try:
                self.json_response(200, rarity_edits.operation_status(WORKBOOK, BACKUPS, path.rsplit('/', 1)[1]))
            except workbook_edits.EditError as error:
                self.json_response(error.status, {'error': str(error)})
            except OSError:
                self.json_response(423, {'error': 'Spreadsheet unavailable; save outcome could not be checked.'})
            return
        if path.startswith('/api/tier-operations/'):
            try:
                self.json_response(200, tier_edits.operation_status(WORKBOOK, BACKUPS, path.rsplit('/',1)[1]))
            except workbook_edits.EditError as error:
                self.json_response(error.status, {'error':str(error)})
            except OSError:
                self.json_response(423, {'error':'Spreadsheet unavailable; save outcome could not be checked.'})
            return
        if path.startswith('/api/category-operations/'):
            try:
                self.json_response(200, category_edits.operation_status(WORKBOOK, BACKUPS, path.rsplit('/', 1)[1]))
            except workbook_edits.EditError as error:
                self.json_response(error.status, {"error": str(error)})
            except OSError:
                self.json_response(423, {"error": "Spreadsheet unavailable; save outcome could not be checked."})
            return
        if path in ("/api/catalog", "/api/revision"):
            try:
                data = ({"sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(), **review_metadata()}
                        if path == "/api/revision" else catalog())
                payload = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
            except Exception as error:
                self.send_error(500, f"Workbook could not be read: {type(error).__name__}")
                return
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
        elif path == "/material-reference.jpg":
            if not REFERENCE.exists():
                self.send_error(404)
                return
            payload = REFERENCE.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "image/jpeg")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
        else:
            super().do_GET()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8087)
    args = parser.parse_args()
    server = ThreadingHTTPServer(("127.0.0.1", args.port), partial(Handler, directory=str(WEB)))
    print(f"Inventory visualizer: http://127.0.0.1:{args.port}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
