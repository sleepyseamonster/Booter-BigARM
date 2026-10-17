"""Rename persistence checks use disposable workbooks; the live catalog is read-only."""
import json
import tempfile
import threading
import unittest
from functools import partial
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from uuid import uuid4

import openpyxl
import name_edits as edits
import server
from test_category_moves import fixture


class NameTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=server.ROOT / 'Logs')
        self.book = Path(self.temp.name) / 'fixture.xlsx'
        self.backups = Path(self.temp.name) / 'backups'
        fixture(self.book)
        book = openpyxl.load_workbook(self.book)
        for sheet, mappings in edits.NAME_LINKS.items():
            sh = book[sheet] if sheet in book.sheetnames else book.create_sheet(sheet)
            for i, (key, label) in enumerate(mappings):
                # Fixtures intentionally include duplicate references to one item.
                column = 20 + i * 2
                sh.cell(5, column, key)
                sh.cell(5, column + 1, label)
                for row in (8, 9):
                    sh.cell(row, column, 'test.guard')
                    sh.cell(row, column + 1, '=Items!B6')
        book['Variants']['B6'] = 'Original authored name'
        book.save(self.book)
        book.close()

    def tearDown(self):
        self.temp.cleanup()

    def body(self, **overrides):
        return {**dict(itemId='test.guard', fromName='Test guard', toName='Guard & <shield>',
                       sourceHash=edits.digest(self.book.read_bytes()), operationId=str(uuid4())), **overrides}

    def test_names_linked_labels_search_and_stable_identity(self):
        original = self.book.read_bytes()
        candidate, receipt = edits.build_name(original, 'test.guard', '  Guard & <shield>  ', 'Test guard')
        values = edits.parse_sheets(candidate)
        self.assertEqual(values['Items'][5][0], 'test.guard')
        self.assertEqual(values['Items'][5][1], 'Guard & <shield>')
        self.assertEqual(values['Items'][5][21], 'Test guard')
        self.assertIn('guard & <shield>', values['Items'][5][30])
        self.assertIn('test guard', values['Items'][5][30])
        self.assertEqual(values['Items'][6][1], 'Test tool')
        self.assertEqual(values['Variants'][5][1], 'Original authored name')
        for sheet, mappings in edits.NAME_LINKS.items():
            headers = values[sheet][4]
            for key, label in mappings:
                for row in values[sheet][5:]:
                    if row[headers.index(key)] == 'test.guard':
                        self.assertEqual(row[headers.index(label)], 'Guard & <shield>')
        from io import BytesIO
        formulas = openpyxl.load_workbook(BytesIO(candidate), data_only=False)
        self.assertEqual(formulas['Item relationships']['U8'].value, '=Items!B6')
        self.assertTrue(formulas['Items']['AE6'].value.startswith('=LOWER('))
        formulas.close()
        self.assertEqual(receipt['toName'], 'Guard & <shield>')
        self.assertEqual(self.book.read_bytes(), original)
        # Leading '=' is literal text, never executable spreadsheet content.
        candidate, _ = edits.build_name(original, 'test.guard', '=Harmless name', 'Test guard')
        self.assertEqual(edits.parse_sheets(candidate)['Items'][5][1], '=Harmless name')

    def test_validation_revision_backup_retry_and_noop(self):
        original = self.book.read_bytes()
        for overrides in [{'toName': ''}, {'toName': '   '}, {'toName': 'a' * 201},
                          {'toName': 'Bad\nname'}, {'toName': '\x00Bad'}, {'toName': None},
                          {'fromName': 'stale name'}, {'sourceHash': 'stale'},
                          {'itemId': 'missing'}, {'operationId': '../bad'}]:
            with self.subTest(overrides=overrides), self.assertRaises(edits.EditError):
                edits.update(self.book, self.backups, self.body(**overrides))
            self.assertEqual(self.book.read_bytes(), original)
        noop = edits.update(self.book, self.backups, self.body(toName='Test guard'))
        self.assertTrue(noop['noop'])
        body = self.body(toName='  Guard & <shield>  ')
        receipt = edits.update(self.book, self.backups, body)
        self.assertEqual(edits.update(self.book, self.backups, body), receipt)
        self.assertEqual((self.backups / 'Name' / (edits.digest(original) + '.xlsx')).read_bytes(), original)
        with self.assertRaises(edits.EditError):
            edits.update(self.book, self.backups, {**body, 'toName': 'Other'})

    def test_locked_workbook_and_lost_receipt(self):
        original = self.book.read_bytes()
        real_write = edits.atomic_write
        body = self.body()

        def locked(path, *args, **kwargs):
            if path == self.book:
                raise PermissionError('Excel lock')
            return real_write(path, *args, **kwargs)

        with patch.object(edits, 'atomic_write', side_effect=locked), self.assertRaises(PermissionError):
            edits.update(self.book, self.backups, body)
        self.assertEqual(self.book.read_bytes(), original)
        self.assertEqual(edits.operation_status(self.book, self.backups, body['operationId'])['status'], 'not-applied')

        def interrupted(path, *args, **kwargs):
            if path.name.endswith('.committed.json'):
                raise OSError('interrupted receipt')
            return real_write(path, *args, **kwargs)

        with patch.object(edits, 'atomic_write', side_effect=interrupted):
            edits.update(self.book, self.backups, body)
        self.assertEqual(edits.operation_status(self.book, self.backups, body['operationId'])['status'], 'committed')

    def test_formula_name_rejected(self):
        book = openpyxl.load_workbook(self.book)
        book['Items']['B6'] = '="Name"'
        book.save(self.book)
        book.close()
        with self.assertRaises(edits.EditError):
            edits.build_name(self.book.read_bytes(), 'test.guard', 'New name', None)

    def test_real_workbook_format_in_memory_only(self):
        original = server.WORKBOOK.read_bytes()
        number, item = edits.master_items(edits.parse_sheets(original))[0]
        candidate, _ = edits.build_name(original, item[0], item[1] + ' check', item[1])
        self.assertEqual(edits.parse_sheets(candidate)['Items'][number - 1][1], item[1] + ' check')

    def test_http_security_save_catalog_and_operation_status(self):
        with patch.object(server, 'WORKBOOK', self.book), patch.object(server, 'BACKUPS', self.backups):
            http = server.ThreadingHTTPServer(('127.0.0.1', 0), partial(server.Handler, directory=str(server.WEB)))
            worker = threading.Thread(target=http.serve_forever, daemon=True)
            worker.start()
            base = f'http://127.0.0.1:{http.server_port}'

            def post(body, origin=base, token=server.TOKEN):
                return urlopen(Request(base + '/api/items/name', data=json.dumps(body).encode(), headers={
                    'Content-Type': 'application/json', 'Origin': origin, 'X-Inventory-Token': token}))

            try:
                for origin, token in [('http://other.test', server.TOKEN), (base, 'bad')]:
                    with self.assertRaises(HTTPError) as error:
                        post(self.body(), origin, token)
                    self.assertEqual(error.exception.code, 403)
                with self.assertRaises(HTTPError) as error:
                    post(self.body(toName=''))
                self.assertEqual(error.exception.code, 400)
                body = self.body()
                with post(body) as response:
                    self.assertEqual(json.load(response)['toName'], 'Guard & <shield>')
                with urlopen(base + '/api/catalog') as response:
                    catalog = json.load(response)
                    self.assertTrue(catalog['nameEditing'])
                    self.assertEqual(catalog['items'][0]['Item name'], 'Guard & <shield>')
                with urlopen(base + '/api/name-operations/' + body['operationId']) as response:
                    self.assertEqual(json.load(response)['status'], 'committed')
                with self.assertRaises(HTTPError) as error:
                    post(self.body(sourceHash='stale'))
                self.assertEqual(error.exception.code, 409)
            finally:
                http.shutdown()
                http.server_close()
                worker.join()


if __name__ == '__main__':
    unittest.main()
