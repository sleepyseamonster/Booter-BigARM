"""Append Canyon Crew and register its explicitly authorized cloth material."""
import json
import re
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import add_arc_power_armor as batch
from add_ancient_salvage import build as append_items
from category_edits import parts_for, patch_cells
from workbook_edits import digest

SPEC = 'Docs/Design/Items/CANYON_CREW_CLOTHING_2026-10-16.json'
OUT = batch.append.ROOT / 'Logs/CanyonCrew20261016'


def register_cloth(snapshot):
    """Append taxonomy without shifting the fixed ranges for other fields."""
    book = openpyxl.load_workbook(BytesIO(snapshot), data_only=False)
    sheet = book['Taxonomy']
    assert sheet.tables['ItemTaxonomy'].ref == 'A5:C242'
    choices = [r[1].value for r in sheet if r[0].value == 'Material family']
    assert 'Cloth' not in choices and len(choices) == len(set(choices))
    choices.append('Cloth')
    cells = {'A243': 'Material family', 'B243': 'Cloth', 'C243': 'Author-authorized Canyon Crew cloth clothing, October 16, 2026. Composition and provenance unresolved; this does not establish ordinary cloth elsewhere.', 'E5': 'Material family choices'}
    cells.update({f'E{n}': value for n, value in enumerate(choices, 6)})
    assert all(sheet[a].value is None for a in cells)
    with ZipFile(BytesIO(snapshot)) as source:
        parts = parts_for(source)
        raw = patch_cells(source.read(parts['Taxonomy']).decode(), cells)
        raw = re.sub(r'(<(?:\w+:)?dimension\b[^>]*ref=")[^"]+', r'\g<1>A1:E243', raw, count=1)
        changed = {parts['Taxonomy']: raw.encode()}
        raw = source.read(parts['Items']).decode()
        old = "'Taxonomy'!$B$159:$B$194"
        assert raw.count(old) == 1
        changed[parts['Items']] = raw.replace(old, f"'Taxonomy'!$E$6:$E${len(choices)+5}").encode()
        for info in source.infolist():
            if info.filename.startswith('xl/tables/'):
                raw = source.read(info).decode()
                if 'name="ItemTaxonomy"' in raw:
                    changed[info.filename] = raw.replace('ref="A5:C242"', 'ref="A5:C243"').encode()
        result = BytesIO()
        with ZipFile(result, 'w') as target:
            for info in source.infolist():
                target.writestr(info, changed.get(info.filename, source.read(info)))
    book.close()
    return result.getvalue(), cells, choices


def build(snapshot, spec):
    registered, cells, choices = register_cloth(snapshot)
    authored = openpyxl.load_workbook(OUT / 'authored.xlsx', read_only=True, data_only=True)
    assert [r[0] for r in list(authored['Material choices'].values)[1:]] == choices
    authored.close()
    result, receipt = append_items(registered, spec)
    before = openpyxl.load_workbook(BytesIO(snapshot), data_only=False)
    after = openpyxl.load_workbook(BytesIO(result), data_only=False)
    for sheet in before:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    other = after[sheet.title][cell.coordinate]
                    assert cell.value == other.value and cell.style_id == other.style_id, (sheet.title, cell.coordinate)
    assert all(after['Taxonomy'][a].value == value for a, value in cells.items())
    validation = next(d for d in after['Items'].data_validations.dataValidation if str(d.sqref) == 'AB6:AB1005')
    assert validation.formula1 == f"'Taxonomy'!$E$6:$E${len(choices)+5}"
    with ZipFile(BytesIO(snapshot)) as old, ZipFile(BytesIO(result)) as new:
        assert old.read('xl/styles.xml') == new.read('xl/styles.xml')
    before.close(); after.close()
    receipt.update(sourceHash=digest(snapshot), resultHash=digest(result), materialRegistry='Cloth added to Taxonomy and Items material dropdown; previous choices retained')
    return result, receipt


def prepare():
    batch.prepare()
    _, _, choices = register_cloth((OUT / 'before.xlsx').read_bytes())
    (OUT / 'material-choices.json').write_text(json.dumps(choices), encoding='utf-8')


if __name__ == '__main__':
    batch.OUT = OUT; batch.SPEC = SPEC; batch.append.build = build
    batch.guard.OUT = OUT; batch.guard.prepare = prepare; batch.guard.apply = batch.apply
    batch.guard.run()
