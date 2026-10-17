"""Rarity persistence and HTTP checks on disposable workbooks only."""
import json
import tempfile
import threading
import unittest
from functools import partial
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from uuid import uuid4

import openpyxl
import rarity_edits as edits
import server
from test_category_moves import fixture


class RarityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=server.ROOT / 'Logs')
        self.book = Path(self.temp.name) / 'fixture.xlsx'
        self.backups = Path(self.temp.name) / 'backups'
        fixture(self.book)
        book = openpyxl.load_workbook(self.book)
        book['Items']['P7'] = 'Extremely rare'
        book.save(self.book)
        book.close()

    def tearDown(self):
        self.temp.cleanup()

    def body(self, **overrides):
        body = dict(itemId='test.guard', fromRarity='', toRarity='Epic',
                    sourceHash=edits.digest(self.book.read_bytes()), operationId=str(uuid4()))
        return {**body, **overrides}

    def test_all_choices_clear_and_legacy_values_preserve_other_cells(self):
        original = self.book.read_bytes()
        for rarity in edits.RARITIES:
            with self.subTest(rarity=rarity):
                candidate, receipt = edits.build_rarity(original, 'test.guard', rarity, '')
                sheets = edits.parse_sheets(candidate)
                self.assertEqual(sheets['Items'][5][15], rarity)
                self.assertEqual(sheets['Items'][6][15], 'Extremely rare')
                self.assertEqual(sheets['Items'][5][2], 'T1')
                self.assertEqual(receipt['changedCells'], {'Items': 1})
                cleared, _ = edits.build_rarity(candidate, 'test.guard', '', rarity)
                self.assertFalse(edits.parse_sheets(cleared)['Items'][5][15])
        candidate, _ = edits.build_rarity(original, 'test.tool', 'Legendary', 'Extremely rare')
        self.assertEqual(edits.parse_sheets(candidate)['Items'][6][15], 'Legendary')
        self.assertEqual(self.book.read_bytes(), original)

    def test_revision_idempotency_backups_invalid_missing_and_noop(self):
        original = self.book.read_bytes()
        for overrides in [{'toRarity': 'Mythic'}, {'toRarity': None}, {'fromRarity': 'Rare'},
                          {'sourceHash': 'stale'}, {'itemId': 'missing'}, {'operationId': '../bad'}]:
            with self.subTest(overrides=overrides), self.assertRaises(edits.EditError):
                edits.update(self.book, self.backups, self.body(**overrides))
            self.assertEqual(self.book.read_bytes(), original)
        noop = edits.update(self.book, self.backups, self.body(toRarity=''))
        self.assertTrue(noop['noop'])
        body = self.body()
        receipt = edits.update(self.book, self.backups, body)
        self.assertEqual(edits.update(self.book, self.backups, body), receipt)
        self.assertEqual((self.backups / 'Rarity' / (edits.digest(original) + '.xlsx')).read_bytes(), original)
        with self.assertRaises(edits.EditError):
            edits.update(self.book, self.backups, {**body, 'toRarity': 'Rare'})

    def test_locked_workbook_interrupted_journal_and_external_writer(self):
        original = self.book.read_bytes()
        real_write = edits.atomic_write
        body = self.body()

        def locked(path, *args, **kwargs):
            if path == self.book:
                raise PermissionError('Excel lock')
            return real_write(path, *args, **kwargs)

        with patch.object(edits, 'atomic_write', side_effect=locked), self.assertRaises(PermissionError):
            edits.update(self.book, self.backups, body)
        self.assertEqual(self.book.read_bytes(), original)
        self.assertEqual(edits.operation_status(self.book, self.backups, body['operationId'])['status'], 'not-applied')

        def interrupted(path, *args, **kwargs):
            if path.name.endswith('.committed.json'):
                raise OSError('journal interrupted')
            return real_write(path, *args, **kwargs)

        with patch.object(edits, 'atomic_write', side_effect=interrupted):
            edits.update(self.book, self.backups, body)
        self.assertEqual(edits.operation_status(self.book, self.backups, body['operationId'])['status'], 'committed')
        current = self.book.read_bytes()

        def conflict(path, *args, **kwargs):
            if path == self.book:
                path.write_bytes(current + b' ')
            return real_write(path, *args, **kwargs)

        with patch.object(edits, 'atomic_write', side_effect=conflict), self.assertRaises(edits.EditError):
            edits.update(self.book, self.backups, self.body(fromRarity='Epic', toRarity='Rare'))
        self.assertEqual(self.book.read_bytes(), current + b' ')

    def test_formula_and_duplicate_ids_rejected(self):
        book = openpyxl.load_workbook(self.book)
        book['Items']['P6'] = '="Rare"'
        book.save(self.book)
        book.close()
        with self.assertRaises(edits.EditError):
            edits.build_rarity(self.book.read_bytes(), 'test.guard', 'Common', '')
        book = openpyxl.load_workbook(self.book)
        book['Items']['P6'] = None
        book['Items']['A7'] = 'test.guard'
        book.save(self.book)
        book.close()
        with self.assertRaises(edits.EditError):
            edits.build_rarity(self.book.read_bytes(), 'test.guard', 'Common', '')

    def test_actual_workbook_format_in_memory_only(self):
        original = server.WORKBOOK.read_bytes()
        sheets = edits.parse_sheets(original)
        number, item = edits.master_items(sheets)[0]
        rarity_column = sheets['Items'][4].index('Rarity')
        old = item[rarity_column] or ''
        destination = 'Legendary' if old != 'Legendary' else 'Junk'
        candidate, receipt = edits.build_rarity(original, item[0], destination, old)
        self.assertEqual(edits.parse_sheets(candidate)['Items'][number - 1][rarity_column], destination)
        self.assertEqual(receipt['changedCells'], {'Items': 1})
        self.assertEqual(server.WORKBOOK.read_bytes(), original)

    def test_http_origin_token_invalid_save_catalog_and_status(self):
        with patch.object(server, 'WORKBOOK', self.book), patch.object(server, 'BACKUPS', self.backups):
            http = server.ThreadingHTTPServer(('127.0.0.1', 0), partial(server.Handler, directory=str(server.WEB)))
            worker = threading.Thread(target=http.serve_forever, daemon=True)
            worker.start()
            base = f'http://127.0.0.1:{http.server_port}'

            def post(body, origin=base, token=server.TOKEN):
                return urlopen(Request(base + '/api/items/rarity', data=json.dumps(body).encode(), headers={
                    'Content-Type': 'application/json', 'Origin': origin, 'X-Inventory-Token': token}))

            try:
                for origin, token in [('http://other.test', server.TOKEN), (base, 'wrong')]:
                    with self.assertRaises(HTTPError) as error:
                        post(self.body(), origin, token)
                    self.assertEqual(error.exception.code, 403)
                for overrides in [{'toRarity': None}, {'toRarity': 'Mythic'}]:
                    with self.assertRaises(HTTPError) as error:
                        post(self.body(**overrides))
                    self.assertEqual(error.exception.code, 400)
                body = self.body()
                with post(body) as response:
                    self.assertEqual(json.load(response)['toRarity'], 'Epic')
                with post(body) as response:
                    self.assertEqual(json.load(response)['operationId'], body['operationId'])
                with urlopen(base + '/api/catalog') as response:
                    catalog = json.load(response)
                    self.assertEqual(catalog['rarityChoices'], edits.RARITIES)
                    self.assertTrue(catalog['rarityEditing'])
                    self.assertEqual(catalog['items'][0]['Rarity'], 'Epic')
                with urlopen(base + '/api/rarity-operations/' + body['operationId']) as response:
                    self.assertEqual(json.load(response)['status'], 'committed')
                with self.assertRaises(HTTPError) as error:
                    post(self.body(sourceHash='stale'))
                self.assertEqual(error.exception.code, 409)
            finally:
                http.shutdown()
                http.server_close()
                worker.join()


if __name__ == '__main__':
    unittest.main()
