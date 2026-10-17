"""Rename the salvage category and append verified expedition recovery concepts."""
import json
from io import BytesIO
from zipfile import ZipFile

import openpyxl
from add_ancient_salvage import ROOT, SOURCE, build as append_items
from category_edits import DEPENDENCIES, parts_for, patch_cells
from split_equipment import verify
from workbook_edits import atomic_write, digest

OLD, NEW = 'Salvage feedstock', 'Salvage'
DESCRIPTION = 'Recovered scrap, discarded supplies and damaged or potentially reusable expedition gear from settlers, Booters and abandoned camps. Inspection and qualification precede reuse; finding an item grants no contents or repair yield.'
DATA = ROOT / 'Docs/Design/Items/EXPEDITION_SALVAGE_PROPOSALS.json'
OUT = ROOT / 'Logs/ExpeditionSalvage'


def rename(snapshot):
    values = openpyxl.load_workbook(BytesIO(snapshot), data_only=True)
    formulas = openpyxl.load_workbook(BytesIO(snapshot), data_only=False)
    changes = {}
    master = [(r[0], n) for n, r in enumerate(values['Items'].values, 1) if len(r)>3 and r[0] and r[3]==OLD]
    assert master, 'Old salvage category is absent; do not run this migration twice.'
    ids = {key for key, _ in master}
    def change(sheet, address, value):changes.setdefault(sheet, {})[address] = value
    for key, n in master:
        change('Items', f'D{n}', NEW)
        change('Items', f'AE{n}', values['Items'][f'AE{n}'].value.replace('; salvage feedstock;', '; salvage;'))
    for name, mappings in DEPENDENCIES.items():
        sheet = values[name];rows = list(sheet.values);header = rows[4]
        for key, category in mappings:
            if key not in header or category not in header:continue
            ki, ci = header.index(key), header.index(category)
            for n, row in enumerate(rows[5:], 6):
                if len(row)>ki and row[ki] in ids:change(name, f'{openpyxl.utils.get_column_letter(ci+1)}{n}', NEW)
    for row in values['Categories']:
        if row[0].value==OLD:
            change('Categories', row[0].coordinate, NEW);change('Categories', row[2].coordinate, DESCRIPTION)
        # The category editor's contiguous dropdown list, if already installed.
        if len(row)>6 and row[6].value==OLD:change('Categories', row[6].coordinate, NEW)
    for row in values['Taxonomy']:
        if row[0].value=='Category' and row[1].value==OLD:
            change('Taxonomy', row[1].coordinate, NEW);change('Taxonomy', row[2].coordinate, DESCRIPTION)
    for row in values['Guide']:
        for cell in row:
            if isinstance(cell.value,str) and OLD in cell.value:change('Guide', cell.coordinate, cell.value.replace(OLD,NEW))
    with ZipFile(BytesIO(snapshot)) as archive:
        parts = parts_for(archive)
        replacements = {parts[name]:patch_cells(archive.read(parts[name]).decode('utf-8'), cells).encode('utf-8') for name,cells in changes.items()}
        stream = BytesIO()
        with ZipFile(stream,'w') as result:
            for member in archive.infolist():result.writestr(member,replacements.get(member.filename,archive.read(member.filename)))
    candidate = stream.getvalue()
    plan = {'changes':[{'sheet':name,'address':address,'value':value,'cachedOnly':formulas[name][address].data_type=='f'} for name,cells in changes.items() for address,value in cells.items()]}
    verify(candidate,snapshot,plan)
    values.close();formulas.close()
    return candidate, sorted(ids)


def build(snapshot):
    renamed, ids = rename(snapshot)
    candidate, receipt = append_items(renamed,json.loads(DATA.read_text(encoding='utf-8')))
    receipt.update(sourceHash=digest(snapshot),renamedFrom=OLD,renamedTo=NEW,reclassifiedIds=ids)
    return candidate, receipt


if __name__ == '__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    original = SOURCE.read_bytes();(OUT/'before.xlsx').write_bytes(original)
    candidate, receipt = build(original)
    (OUT/'after.xlsx').write_bytes(candidate)
    atomic_write(SOURCE,candidate,receipt['sourceHash'])
    (OUT/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
    print(json.dumps(receipt,indent=2))
