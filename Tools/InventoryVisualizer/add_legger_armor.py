"""Prepare and merge three Artifact Tool-authored armor records without re-exporting the master.

Full-master imports previously exhausted Artifact Tool's heap; a small authored
record workbook avoids that limitation and preserves all unrelated OOXML parts.
"""
import argparse
import json
import re
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

import openpyxl
from openpyxl.formula.translate import Translator
from category_edits import CELL, parts_for, patch_cells
from workbook_edits import atomic_write, digest

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'Docs/Design/Items/MASTER_ITEM_DATABASE.xlsx'
SPEC = ROOT / 'Docs/Design/Items/LEGGER_ARMOR_PROPOSALS_2026-10-15.json'
OUT = ROOT / 'Logs/LeggerArmorThree20261015'


def records_for(headers):
    spec = json.loads(SPEC.read_text(encoding='utf-8'))
    result = []
    for item in spec['items']:
        fields = {
            'Item ID': item['id'], 'Item name': item['name'], 'Tier': 'Not applicable',
            'Category': 'Legger armor', 'Subtype': 'Legger armor', 'Implementation scope': 'New proposal',
            'Purpose': item['purpose'], 'Equipment slot': 'Legger fitment unresolved',
            'Carrier metadata': 'Legger attachment / replacement intended; physical cargo policy unchanged',
            'Acquisition': 'Proposed contemporary fabrication or specialist trade; availability unresolved.',
            'Crafting / processing': 'Candidate finished armor using existing material families. No recipe, input quantities, production time or yield authored.',
            'Supported behavior': 'Not implemented', 'Limitations': item['limitations'],
            'Open design decisions': 'Artwork direction: ' + item['icon'] + ' Fit, mass, fastening, removal, repair and performance remain proposals.',
            'Search aliases': item['aliases'], 'Sources': spec['source'], 'Item type': 'Armor',
            'Processing stage': 'Finished item', 'Origin family': 'Manufactured / unresolved',
            'Primary use': 'Protection', 'Material family': item['material'], 'Physical form': item['form'],
            'Design status': 'Proposal', 'Classification basis': 'Author-requested Legger armor expansion. Intended coverage only; does not establish a working equipment slot or protection behavior.',
            'Tier basis': 'Finished endpoint; material tier is not applicable. Input tiers require an authored recipe.',
            'Finished-item input tier': 'Unassigned', 'Crafting role': 'Finished item'
        }
        result.append([fields.get(header) for header in headers])
    return result


def build(snapshot, records):
    before = openpyxl.load_workbook(BytesIO(snapshot), data_only=False)
    cached = openpyxl.load_workbook(BytesIO(snapshot), data_only=True)
    headers = list(cached['Items'].values)[4][:40]
    old_items = [r[:40] for r in list(cached['Items'].values)[5:] if r[0]]
    ids = [r[0] for r in old_items]
    assert len(ids) == len(set(ids))
    assert len(records) == 3 and len({r[0] for r in records}) == 3
    assert not set(ids) & {r[0] for r in records}, 'Proposal already exists'
    assert not {r[1].casefold() for r in old_items} & {r[1].casefold() for r in records}
    allowed = {}
    for row in cached['Taxonomy'].values:
        if len(row) > 1 and row[0] and row[1]:
            allowed.setdefault(row[0], set()).add(row[1])
    for row in records:
        for key, value in zip(headers, row):
            if key in allowed and value is not None:
                assert value in allowed[key], (key, value)
    template = next(n for n, r in enumerate(cached['Items'].values, 1) if r[0] == 'proposal.legger.shell_guard')
    last = int(re.search(r'\d+$', before['Items'].tables['ItemsTable'].ref).group())
    assert last + 3 <= 1005
    changes, expected, table_ends = {}, {}, {}
    with ZipFile(BytesIO(snapshot)) as archive:
        parts = parts_for(archive)

        def append(sheet, template_row, number, values, caches):
            raw = changes.get(parts[sheet], archive.read(parts[sheet])).decode('utf-8')
            prefix = re.search(r'<((?:\w+:)?)worksheet\b', raw).group(1)
            model = re.search(rf'<(?:\w+:)?row\b[^>]*\br="{template_row}"[^>]*>.*?</(?:\w+:)?row>', raw, re.S).group()
            styles = {re.sub(r'\d+', '', m.group(1)): m.group() for m in re.finditer(CELL, model, re.S)}
            cells = []
            for column, value in enumerate(values, 1):
                letter = openpyxl.utils.get_column_letter(column)
                cell = styles.get(letter, f'<{prefix}c r="{letter}{template_row}"/>')
                opening = re.sub(r'\br="[A-Z]+\d+"', f'r="{letter}{number}"', cell[:cell.index('>')].rstrip('/'))
                opening = re.sub(r'\s+t="[^"]*"', '', opening)
                if value is None:
                    cells.append(opening + '/>')
                elif isinstance(value, str) and value.startswith('='):
                    cache = caches[column]
                    cells.append(opening + (' t="str"' if isinstance(cache, str) else '') + f'><{prefix}f>{escape(value[1:])}</{prefix}f><{prefix}v>{escape(str(cache))}</{prefix}v></{prefix}c>')
                else:
                    cells.append(opening + f' t="inlineStr"><{prefix}is><{prefix}t xml:space="preserve">{escape(str(value))}</{prefix}t></{prefix}is></{prefix}c>')
            content = f'<{prefix}row r="{number}">' + ''.join(cells) + f'</{prefix}row>'
            pattern = rf'<(?:\w+:)?row\b[^>]*\br="{number}"[^>]*(?:/>|>.*?</(?:\w+:)?row>)'
            existing = re.search(pattern, raw, re.S)
            if existing:
                assert not re.search(r'<(?:\w+:)?[vf]|<(?:\w+:)?is', existing.group()), 'New row contains data'
                raw = raw[:existing.start()] + content + raw[existing.end():]
            else:
                raw = raw.replace(f'</{prefix}sheetData>', content + f'</{prefix}sheetData>')
            raw = re.sub(r'(<(?:\w+:)?dimension\b[^>]*\bref="[^":]+:[A-Z]+)(\d+)', lambda m: m[1] + str(max(int(m[2]), number)), raw)
            changes[parts[sheet]] = raw.encode('utf-8')
            expected[(sheet, number)] = (values, caches)

        for offset, record in enumerate(records, 1):
            n = last + offset
            values = list(record)
            search = '; '.join(str(values[c] or '') for c in [0,1,3,4,23,24,25,26,27,28,29,21,2,33,34]).lower()
            caches = {31: search,36:0,37:0,38:0,39:'No production route listed',40:99}
            for col in caches:
                letter = openpyxl.utils.get_column_letter(col)
                values[col-1] = Translator(before['Items'].cell(template, col).value, origin=f'{letter}{template}').translate_formula(f'{letter}{n}')
            append('Items', template, n, values, caches)
            u = before['Item uses'].max_row + offset
            append('Item uses', before['Item uses'].max_row, u, [record[0],f'=Items!B{n}',f'=Items!D{n}','Primary',f'=Items!AA{n}','Author-requested armor concept; intended use only.'], {2:record[1],3:'Legger armor',5:'Protection'})
            append('Craft audit', template, n, [record[0],f'=Items!B{n}',f'=Items!X{n}',f'=Items!C{n}','Finished armor proposal',f'=Items!AJ{n}','No recipe established','Fit, facilities and input quantities unresolved.','Proposal; not qualified for crafting',record[20]], {2:record[1],3:'Armor',4:'Not applicable',6:99})
        table_ends = {'ItemsTable':last+3,'ItemUseRelationships':before['Item uses'].max_row+3,'CraftAudit':last+3}
        for name in archive.namelist():
            if name.startswith('xl/tables/') and name.endswith('.xml'):
                raw = archive.read(name).decode('utf-8')
                for table, end in table_ends.items():
                    if f'name="{table}"' in raw:
                        old = re.search(r'\bref="([^"]+)"',raw)[1]
                        changes[name] = raw.replace(f'ref="{old}"',f'ref="{re.sub(r"\d+$",str(end),old)}"').encode('utf-8')
        all_items = old_items + records
        summary = {}
        for row in before['Categories']:
            if row[0].value == 'Legger armor' and row[4].value == 'Primary category':
                summary[row[1].coordinate] = sum(r[3]=='Legger armor' for r in all_items)
        changes[parts['Categories']] = patch_cells(archive.read(parts['Categories']).decode('utf-8'),summary).encode('utf-8')
        for row in before['Collection targets']:
            if row[1].value == 'Armor':
                excluded = set(re.findall(r"'?Items'?!\$A\$6:\$A\$\d+,\"([^\"]+)\"", row[3].value or ''))
                count = sum(r[23]=='Armor' and r[29]!='Withheld assumption' and r[0] not in excluded for r in all_items)
                caches = {row[3].coordinate:count,row[4].coordinate:max(0,row[2].value-count)}
                changes[parts['Collection targets']] = patch_cells(archive.read(parts['Collection targets']).decode('utf-8'),caches).encode('utf-8')
        for sheet, part in parts.items():
            raw = changes.get(part,archive.read(part)).decode('utf-8')
            updated = raw.replace(f'$A$6:$A${last}',f'$A$6:$A${last+3}')
            if raw != updated:
                changes[part] = updated.encode('utf-8')
        out = BytesIO()
        with ZipFile(out,'w') as target:
            for info in archive.infolist():
                target.writestr(info,changes.get(info.filename,archive.read(info.filename)))
    candidate = out.getvalue()
    after = openpyxl.load_workbook(BytesIO(candidate),data_only=False)
    values = openpyxl.load_workbook(BytesIO(candidate),data_only=True)
    for sheet in before:
        for row in sheet:
            for cell in row:
                if (sheet.title,cell.row) not in expected:
                    assert after[sheet.title][cell.coordinate].value == cell.value,(sheet.title,cell.coordinate)
    for (sheet,n),(row,caches) in expected.items():
        for col,value in enumerate(row,1):
            assert after[sheet].cell(n,col).value == value,(sheet,n,col)
            assert values[sheet].cell(n,col).value == caches.get(col,value),(sheet,n,col,'cache')
    assert [r[0] for r in list(values['Items'].values)[5:] if r[0]] == ids + [r[0] for r in records]
    assert next(r[1] for r in values['Categories'].values if r[0]=='Legger armor') == 10
    for book in [before,cached,after,values]:
        book.close()
    return candidate, {'beforeHash':digest(snapshot),'afterHash':digest(candidate),'ids':[r[0] for r in records],'changedParts':sorted(changes),'rows':[last+1,last+3]}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare','apply'])
    args = parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    if args.action == 'prepare':
        snapshot = SOURCE.read_bytes()
        book = openpyxl.load_workbook(BytesIO(snapshot),read_only=True,data_only=True)
        headers = list(book['Items'].values)[4][:40]
        plan = {'sourceHash':digest(snapshot),'headers':headers,'records':records_for(headers)}
        book.close()
        (OUT/'before.xlsx').write_bytes(snapshot)
        (OUT/'plan.json').write_text(json.dumps(plan,indent=2),encoding='utf-8')
        print('Prepared three armor records.')
    else:
        plan = json.loads((OUT/'plan.json').read_text(encoding='utf-8'))
        authored = openpyxl.load_workbook(OUT/'authored.xlsx',read_only=True,data_only=True)
        records = [list(r)[:40] for r in list(authored['Legger armor'].values)[1:4]]
        authored.close()
        assert records == plan['records'], 'Authored values differ from approved proposals'
        candidate,receipt = build((OUT/'before.xlsx').read_bytes(),records)
        atomic_write(SOURCE,candidate,plan['sourceHash'])
        (OUT/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
        print(json.dumps(receipt,indent=2))
