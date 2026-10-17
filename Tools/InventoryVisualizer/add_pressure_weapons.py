"""Refine two existing pressure items and append three family-specific records."""
import gc
import json
import re
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import add_five_clothing_families as refine
from audit_manufactured_stock import NAME_LINKS
from category_edits import parts_for, patch_cells, verify_move
from tier_edits import build_tier
from workbook_edits import atomic_write, digest, parse_sheets

OUT = refine.ROOT / 'Logs/PressureWeapons20261016'
DATA = refine.ROOT / 'Docs/Design/Items/PRESSURE_WEAPONS_2026-10-16.json'


def configure():
    spec = json.loads(DATA.read_bytes())
    refine.OUT = OUT; refine.DATA = DATA
    refine.REUSE = {str(n): key for n, key in enumerate(spec['reusedIds'])}
    refine.specification = lambda: spec
    return spec


def build(snapshot, spec):
    gc.collect()
    configure()
    original = snapshot
    old = {r[0]: r for r in parse_sheets(snapshot)['Items'][5:] if r[0]}
    reused = {i['id']: i for i in spec['items'] if i['id'] in spec['reusedIds']}
    for key, item in reused.items():
        tier = item['fields'].get('Tier', spec['defaults']['Tier'])
        snapshot, _ = build_tier(snapshot, key, tier, old[key][2])
        gc.collect()
    # The older refinement adapter serializes None as an empty string. Keep
    # existing undefined stats untouched; new rows still receive true blanks.
    prepared = json.loads(json.dumps(spec))
    blanks = {k: v for k, v in prepared['defaults'].items() if v is None}
    prepared['defaults'] = {k: v for k, v in prepared['defaults'].items() if v is not None}
    for item in prepared['items']:
        if item['id'] not in reused:
            item['fields'] = {**blanks, **item['fields']}
    candidate, receipt = refine.build(snapshot, prepared)
    gc.collect()
    sheets = parse_sheets(candidate)
    edits = {}
    names = {key: item['name'] for key, item in reused.items()}
    for sheet, mappings in NAME_LINKS.items():
        h = sheets[sheet][4]
        for key_field, name_field in mappings:
            for n, row in enumerate(sheets[sheet][5:], 6):
                if row[h.index(key_field)] in names:
                    edits.setdefault(sheet, {})[f'{openpyxl.utils.get_column_letter(h.index(name_field)+1)}{n}'] = names[row[h.index(key_field)]]
    renamed = {old[key][1]: item['name'] for key, item in reused.items() if old[key][1] != item['name']}
    pattern = re.compile('|'.join(re.escape(name) for name in sorted(renamed, key=len, reverse=True)))
    formulas = openpyxl.load_workbook(BytesIO(candidate), data_only=False)
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                if cell.data_type != 'f': continue
                cached = sheets[sheet.title][cell.row-1][cell.column-1]
                if isinstance(cached, str) and pattern.search(cached):
                    edits.setdefault(sheet.title, {})[cell.coordinate] = pattern.sub(lambda m: renamed[m.group()], cached)
    formulas.close()
    del formulas, sheet, row, cell
    gc.collect()
    output = BytesIO()
    with ZipFile(BytesIO(candidate)) as source, ZipFile(output, 'w') as target:
        parts = parts_for(source)
        changed = {parts[s]: patch_cells(source.read(parts[s]).decode(), cells).encode() for s, cells in edits.items()}
        for info in source.infolist(): target.writestr(info, changed.get(info.filename, source.read(info)))
    result = output.getvalue()
    verify_move(candidate, result, edits)
    receipt.pop('pieceCount', None)
    receipt.update(sourceHash=digest(original), resultHash=digest(result), catalogItemCount=len(spec['items']), renamedItems=renamed,
                   scope='Two existing IDs refined; three new records. Tier/name links refreshed; formulas and unrelated records retained.')
    return result, receipt


def apply():
    plan = json.loads((OUT / 'plan.json').read_bytes())
    authored = openpyxl.load_workbook(OUT / 'authored.xlsx', read_only=True, data_only=True)
    rows = list(authored['Catalog inputs'].values)[1:]
    assert len(rows) == len(plan['spec']['items'])
    for row, item in zip(rows, plan['spec']['items']):
        assert list(row[:6]) == [item[k] for k in ['id', 'name', 'purpose', 'icon', 'material', 'use']]
        assert json.loads(row[6]) == item['fields']
    authored.close()
    candidate, receipt = build((OUT / 'before.xlsx').read_bytes(), plan['spec'])
    (OUT / 'after.xlsx').write_bytes(candidate)
    refine.guard.release_guard()
    atomic_write(refine.append.SOURCE, candidate, plan['sourceHash'])
    (OUT / 'receipt.json').write_text(json.dumps(receipt, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(receipt, indent=2))


if __name__ == '__main__':
    configure()
    refine.guard.OUT = OUT; refine.guard.prepare = refine.prepare; refine.guard.apply = apply
    refine.guard.run()
