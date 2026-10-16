"""Author five clothing families as individual pieces, reusing electrical wear IDs."""
import json
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import append_catalog_proposals as append
import organize_camp_items as guard
from category_edits import parts_for, patch_cells, verify_move
from workbook_edits import atomic_write, digest, parse_sheets

ROOT=append.ROOT
OUT=ROOT/'Logs/FiveClothingFamilies20261015'
DATA=ROOT/'Docs/Design/Items/FIVE_CLOTHING_FAMILIES_2026-10-15.json'
REUSE={'top':'author.electrical.insulated_suit','gloves':'author.electrical.insulated_gloves','boots':'author.electrical.insulated_boots'}
FAMILIES=[
    ('hide','Hide','Hide','T2','T1','Small fitted hide panels, overlapped seams and simple metal closures. Usable hide yields are extremely low.'),
    ('fur','Fur','Fur','T2','T1','Carefully joined fur-bearing hide patches with restrained trim. Usable fur yields are extremely low; no abundant fur supply assumed.'),
    ('fine_chainmail','Delicate chainmail','Metal','T3','T2','Fine small metal rings with a light drape and narrow rolled edges; no textile backing. Fine-ring work is a labor-intensive proposal.'),
    ('insulated','Electrical insulating','Mixed / unresolved','T3','T2','Overlapping qualified husk/hide-derived insulating layers, guarded joins and rigid shell spacers. Specific dielectric materials and service limits remain unresolved.'),
    ('chitin','Chitin','Shell / chitin','T2','T1','Thin overlapping chitin scales, smooth turned edges and small metal joining points. A flexible clothing silhouette rather than heavy armor plating.'),
]
PIECES=[('top','jacket','One upper-body garment; no trousers, gloves, footwear or headwear included.'),
        ('legs','trousers','One lower-body garment; footwear remains separate.'),
        ('boots','boots','One matched pair of footwear; no leggings or trousers included.'),
        ('gloves','gloves','One matched pair of hand coverings; sleeves and upper-body garment remain separate.'),
        ('head','hood','One head covering; face, sight, hearing and personal canister clearance require design.')]


def specification():
    defaults={'Tier':'T2','Category':'Clothing','Subtype':'Clothing family piece','Implementation scope':'New proposal',
        'Equipment slot':'Unassigned; no new slot approved','Carrier metadata':'Booter-worn garment concept; fit, packing and canister clearance unresolved.',
        'Stack limit':None,'Mass (game units)':None,'Acquisition':'Proposed garment craft or trade; availability and price unresolved.',
        'Crafting / processing':'Candidate construction from qualified stock and compatible joining parts. No full-set recipe, quantities, yield or crafting time assigned.',
        'Supported behavior':'Not implemented','Limitations':'Fit, comfort, flexibility, mass, wear and protective performance require design. No wood, ordinary cloth, conventional abundant hide/fur supply, free-water treatment, set bonus or new equipment slots.',
        'Rarity':None,'Trade value':None,'Damage':None,'Protection':None,'Durability':None,
        'Item type':'Wearable','Processing stage':'Finished item','Origin family':'Manufactured / unresolved',
        'Primary use':'Clothing','Material family':'Hide','Physical form':'Assembly','Design status':'Proposal',
        'Classification basis':'Independently listed clothing piece. Family membership is an organizational grouping, not another counted inventory item or a set bonus.',
        'Tier basis':'Proposed finished-garment planning tier; recipe and performance unresolved.',
        'Finished-item input tier':'T1','Crafting role':'Finished item'}
    items=[];groups=[]
    for family,label,material,tier,input_tier,look in FAMILIES:
        keys=[]
        for piece,noun,coverage in PIECES:
            key=REUSE[piece] if family=='insulated' and piece in REUSE else f'proposal.clothing.{family}.{piece}'
            if family=='fine_chainmail' and piece=='head':noun='coif'
            if family=='chitin' and piece=='head':noun='cap'
            name=f'{label} {noun}'
            fields={'Subtype':f'{label} clothing / {piece}','Tier':tier,'Finished-item input tier':input_tier,
                'Material family':material,'Search aliases':f'{family.replace("_"," ")} clothing; {label.lower()} clothing family; {piece}; separate garment',
                'Open design decisions':f'{coverage} Appearance: {look} Joining, sizing, service life and useful performance require design.',
                'Tier basis':f'Proposed {tier} garment assembly using {input_tier} inputs. Quantities, stock qualification and recipe remain unresolved.'}
            if family=='insulated':
                fields.update({'Primary use':'Electrical insulation','Implementation scope':'Author-directed concept',
                    'Limitations':'Electrical insulation is an intended use, not a qualified safety rating or immunity. Fit, isolation, joins, service limits, wear and performance remain unresolved. No ordinary cloth, wood, sealed breathing system, set bonus or new slots.',
                    'Equipment slot':{'top':'Torso / fit unresolved','gloves':'Hands / fit unresolved','boots':'Feet / fit unresolved'}.get(piece,'Unassigned; no new slot approved')})
            purpose=f'{label} clothing family: one separately listed {noun}. {coverage} {look} Functions and performance remain proposals.'
            items.append(dict(id=key,name=name,purpose=purpose,icon=f'{name}. {look} Show only this garment or matched pair; no full outfit or wearer.',material=material,use='Electrical insulation' if family=='insulated' else 'Clothing',fields=fields))
            keys.append(key)
        groups.append(dict(family=family,itemIds=keys,itemCount=5))
    return dict(source='Author requested hide, fur, delicate chainmail, insulated and chitin clothing families as separate individual garments, October 15, 2026. Docs/WORLD_BASIS.md controls no-cloth/no-wood and low-yield hide/fur rules. Existing electrical clothing IDs reused; generic suit refined to one jacket.',
        scope='25 separate clothing pieces across five families: 22 new entries and three reused electrical-wear IDs. No counted full-set item, set bonus, runtime equipment, world identity, streaming or saved gameplay changes.',
        defaults=defaults,reusedIds=list(REUSE.values()),groups=groups,items=items)


def prepare():
    OUT.mkdir(parents=True,exist_ok=True)
    spec=specification()
    if DATA.exists():assert json.loads(DATA.read_text(encoding='utf-8'))==spec
    else:DATA.write_bytes((json.dumps(spec,indent=2)+'\n').encode())
    snapshot=append.SOURCE.read_bytes()
    ids={r[0] for r in parse_sheets(snapshot)['Items'][5:] if r[0]}
    assert set(REUSE.values())<=ids
    (OUT/'before.xlsx').write_bytes(snapshot)
    (OUT/'plan.json').write_text(json.dumps(dict(sourceHash=digest(snapshot),spec=spec),indent=2),encoding='utf-8')


def build(snapshot,spec):
    sheets=parse_sheets(snapshot);headers=sheets['Items'][4]
    reused={i['id']:i for i in spec['items'] if i['id'] in REUSE.values()}
    edits={}
    for key,item in reused.items():
        n,old=next((n,r) for n,r in enumerate(sheets['Items'],1) if r and r[0]==key)
        new=dict(spec['defaults']);new.update(item['fields']);new.update({'Item name':item['name'],'Purpose':item['purpose'],'Sources':spec['source']})
        search={k:old[i] for i,k in enumerate(headers)};search.update(new)
        new['Search terms']='; '.join(str(search.get(k) or '') for k in ['Item ID','Item name','Category','Subtype','Item type','Processing stage','Origin family','Primary use','Material family','Physical form','Design status','Search aliases','Tier','Finished-item input tier','Crafting role']).lower()
        new['Tier sort order']=int(new['Tier'][1:]) if new['Tier'] in ['T0','T1','T2','T3','T4'] else 99
        for field,value in new.items():
            if field in headers:edits.setdefault('Items',{})[f'{openpyxl.utils.get_column_letter(headers.index(field)+1)}{n}']=value if value is not None else ''
        # Keep exact-ID linked metadata and cached formula values consistent.
        for sheet,rows in sheets.items():
            if sheet=='Items' or len(rows)<5:continue
            h=rows[4]
            for number,row in enumerate(rows[5:],6):
                if key not in row:continue
                for ci,field in enumerate(h):
                    if field in ['Item name','Tier','Item type','Processing stage','Origin family','Material family','Primary use','Tier basis','Finished-item input tier','Tier sort order'] and field in new:
                        edits.setdefault(sheet,{})[f'{openpyxl.utils.get_column_letter(ci+1)}{number}']=new[field]
    output=BytesIO()
    with ZipFile(BytesIO(snapshot)) as archive:
        parts=parts_for(archive)
        changed={parts[s]:patch_cells(archive.read(parts[s]).decode(),cells).encode() for s,cells in edits.items()}
        with ZipFile(output,'w') as result:
            for info in archive.infolist():result.writestr(info,changed.get(info.filename,archive.read(info.filename)))
    revised=output.getvalue();verify_move(snapshot,revised,edits)
    additions=dict(spec);additions['items']=[i for i in spec['items'] if i['id'] not in reused]
    candidate,receipt=append.build(revised,additions)
    receipt.update(sourceHash=digest(snapshot),reusedIds=list(reused),pieceCount=len(spec['items']))
    return candidate,receipt


def apply():
    plan=json.loads((OUT/'plan.json').read_text(encoding='utf-8'));spec=plan['spec']
    authored=openpyxl.load_workbook(OUT/'authored.xlsx',read_only=True,data_only=True)
    rows=list(authored['Catalog inputs'].values)[1:];assert len(rows)==len(spec['items'])
    for row,item in zip(rows,spec['items']):
        assert list(row[:6])==[item[k] for k in ['id','name','purpose','icon','material','use']]
        assert json.loads(row[6])==item['fields']
    authored.close()
    candidate,receipt=build((OUT/'before.xlsx').read_bytes(),spec)
    (OUT/'after.xlsx').write_bytes(candidate)
    guard.release_guard();atomic_write(append.SOURCE,candidate,plan['sourceHash'])
    (OUT/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
    update_manifest(spec,plan['sourceHash'],digest(candidate))
    print('Saved',len(spec['items']),'individual clothing pieces; reused IDs:',len(REUSE),'; no full-set item.')


def update_manifest(spec,before_hash,after_hash):
    path=ROOT/'Docs/Design/Items/EQUIPMENT_CLASSIFICATION.json'
    original=path.read_bytes();manifest=json.loads(original)
    reused={i['id']:i for i in spec['items'] if i['id'] in REUSE.values()}
    for item in manifest['items']:
        if item['id'] in reused:
            item.update(name=reused[item['id']]['name'],reason='Individual garment refined within the requested clothing family; existing ID retained.')
    manifest['sourceHashAfter']=after_hash
    manifest.setdefault('revisions',[]).append(dict(reason=spec['scope'],reusedIds=list(reused),sourceHashBefore=before_hash,sourceHashAfter=after_hash))
    atomic_write(path,(json.dumps(manifest,indent=2)+'\n').encode(),digest(original))


if __name__=='__main__':
    guard.OUT=OUT;guard.prepare=prepare;guard.apply=apply;guard.run()
