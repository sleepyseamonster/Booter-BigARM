"""Category writes use temporary fixtures; real-format validation stays in memory."""
import json
import tempfile
import threading
import unittest
from functools import partial
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
import category_edits as edits
import server
import workbook_edits


def fixture(path):
    b = openpyxl.Workbook();s = b.active;s.title = 'Items'
    headers = ['Item ID','Item name','Tier','Category','Subtype','Implementation scope','Purpose','Equipment slot','Carrier metadata','Stack limit','Mass (game units)','Acquisition','Crafting / processing','Supported behavior','Limitations','Rarity','Trade value','Damage','Protection','Durability','Open design decisions','Search aliases','Sources','Item type','Processing stage','Origin family','Primary use','Material family','Physical form','Design status','Search terms','Classification basis','Tier basis','Finished-item input tier','Crafting role']
    s.append(['Fixture']);s.cell(5,1)
    for i,h in enumerate(headers,1):s.cell(5,i,h)
    for row,item,name,cat in [(6,'test.guard','Test guard','Armor'),(7,'test.tool','Test tool','Tools')]:
        s.cell(row,1,item);s.cell(row,2,name);s.cell(row,3,'T1');s.cell(row,4,cat)
        s.cell(row,24,'Finished item');s.cell(row,31,f'=LOWER(A{row}&"; "&B{row}&"; "&D{row}&"; "&E{row}&"; "&X{row}&"; "&Y{row}&"; "&Z{row}&"; "&AA{row}&"; "&AB{row}&"; "&AC{row}&"; "&AD{row}&"; "&V{row}&"; "&C{row}&"; "&AH{row}&"; "&AI{row})')
        s.cell(row,32,'Historical decision');s.cell(row,4).number_format='@'
    v=DataValidation(type='list',formula1="'Taxonomy'!$B$6:$B$7");v.add('D6:D1005');s.add_data_validation(v)
    c=b.create_sheet('Categories')
    for i,h in enumerate(['Category','Current coverage','Existing basis','Source','Grouping'],1):c.cell(5,i,h)
    for n,name in enumerate(['Armor','Tools','Empty category'],6):
        c.cell(n,1,name);c.cell(n,2,f'=COUNTIF(Items!$D$6:$D$1005,A{n})');c.cell(n,5,'Primary category')
    c.cell(9,1,'Overlapping family');c.cell(9,5,'Family')
    for name,keys in edits.DEPENDENCIES.items():
        sh=b.create_sheet(name);cols=[]
        for key,cat in keys:cols.extend([key,cat])
        for i,h in enumerate(cols,1):sh.cell(5,i,h)
        for i,h in enumerate(cols,1):sh.cell(6,i,'test.guard' if 'ID' in h else '=Items!D6')
    b.create_sheet('Variants').cell(6,4,'Equipment')
    b.save(path);b.close()


class CategoryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=server.ROOT/'Logs');self.root=Path(self.temp.name)
        self.book=self.root/'fixture.xlsx';self.backups=self.root/'backups';fixture(self.book)

    def tearDown(self):self.temp.cleanup()

    def body(self,**overrides):
        return dict(itemId='test.guard',fromCategory='Armor',toCategory='Empty category',sourceHash=edits.digest(self.book.read_bytes()),operationId=str(uuid4()),**overrides)

    def test_dependencies_empty_registry_keep_and_idempotency(self):
        original=self.book.read_bytes();body=self.body()
        receipt=edits.move(self.book,self.backups,body)
        self.assertEqual(edits.move(self.book,self.backups,body),receipt)
        with self.assertRaises(edits.EditError):edits.move(self.book,self.backups,{**body,'toCategory':'Tools'})
        w=openpyxl.load_workbook(self.book,data_only=True)
        self.assertEqual(w['Items']['D6'].value,'Empty category');self.assertEqual(w['Items']['A6'].value,'test.guard')
        self.assertIn('empty category',w['Items']['AE6'].value)
        self.assertEqual(w['Items']['AF6'].value,'Historical decision');self.assertEqual(w['Variants']['D6'].value,'Equipment')
        for sn in edits.DEPENDENCIES:
            self.assertEqual(w[sn]['B6'].value,'Empty category')
        self.assertEqual(w['Categories']['B6'].value,0);self.assertEqual(w['Categories']['B8'].value,1)
        self.assertNotIn('Overlapping family',edits.registry(list(w['Categories'].values)))
        w.close()
        self.assertEqual((self.backups/(edits.digest(original)+'.xlsx')).read_bytes(),original)
        formulas=openpyxl.load_workbook(self.book,data_only=False)
        self.assertEqual(formulas['Item uses']['B6'].value,'=Items!D6')
        self.assertIn('InventoryCategoryChoices',formulas.defined_names);formulas.close()
        self.assertEqual(edits.operation_status(self.book,self.backups,body['operationId'])['status'],'committed')

    def test_undo_preserves_intervening_changes_and_rejects_changed_item(self):
        original=self.body();edits.move(self.book,self.backups,original)
        other=dict(itemId='test.tool',fromCategory='Tools',toCategory='Armor',sourceHash=edits.digest(self.book.read_bytes()),operationId=str(uuid4()))
        edits.move(self.book,self.backups,other)
        undo=dict(itemId='test.guard',fromCategory='Empty category',toCategory='Armor',undoOf=original['operationId'],sourceHash=edits.digest(self.book.read_bytes()),operationId=str(uuid4()))
        edits.move(self.book,self.backups,undo)
        values=edits.parse_sheets(self.book.read_bytes())
        self.assertEqual([r[3] for _,r in edits.master_items(values)],['Armor','Armor'])
        again={**undo,'operationId':str(uuid4()),'sourceHash':edits.digest(self.book.read_bytes())}
        with self.assertRaises(edits.EditError):edits.move(self.book,self.backups,again)

    def test_invalid_stale_same_category_unknown_action(self):
        original=self.book.read_bytes()
        for change in [{'toCategory':'All items'},{'sourceHash':'stale'},{'fromCategory':'Tools'},{'operationId':'../../bad'}]:
            body={**self.body(),**change}
            with self.assertRaises(edits.EditError):edits.move(self.book,self.backups,body)
            self.assertEqual(self.book.read_bytes(),original)
        same={**self.body(),'toCategory':'Armor'};self.assertTrue(edits.move(self.book,self.backups,same)['noop'])
        self.assertEqual(self.book.read_bytes(),original)
        with self.assertRaises(edits.EditError):workbook_edits.mutate(self.book,self.root/'review',self.backups,'unexpected','test.guard',edits.digest(original))

    def test_interrupted_commit_reconciliation_and_lock(self):
        body=self.body();real=edits.atomic_write
        def fail_journal(path,*args,**kwargs):
            if path.name.endswith('.committed.json'):raise OSError('lost journal write')
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=fail_journal):edits.move(self.book,self.backups,body)
        self.assertEqual(edits.operation_status(self.book,self.backups,body['operationId'])['status'],'committed')
        next_body={**self.body(),'fromCategory':'Empty category','toCategory':'Tools'}
        before=self.book.read_bytes()
        def locked(path,*args,**kwargs):
            if path==self.book:raise PermissionError('locked')
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=locked):
            with self.assertRaises(PermissionError):edits.move(self.book,self.backups,next_body)
        self.assertEqual(before,self.book.read_bytes())
        self.assertEqual(edits.operation_status(self.book,self.backups,next_body['operationId'])['status'],'not-applied')
        prepared=self.backups/f"category-{next_body['operationId']}.prepared.json"
        prepared_bytes=prepared.read_bytes()
        edits.move(self.book,self.backups,next_body)
        self.assertEqual(prepared.read_bytes(),prepared_bytes)
        next_body={**self.body(),'fromCategory':'Tools','toCategory':'Armor'}
        before=self.book.read_bytes()
        with patch.object(edits,'atomic_write',side_effect=locked):
            with self.assertRaises(PermissionError):edits.move(self.book,self.backups,next_body)
        self.book.write_bytes(before+b' ')
        self.assertEqual(edits.operation_status(self.book,self.backups,next_body['operationId'])['status'],'uncertain')

    def test_external_writer_conflict(self):
        real=edits.atomic_write;body=self.body();external=self.book.read_bytes()+b' '
        def conflict(path,*args,**kwargs):
            if path==self.book:path.write_bytes(external)
            return real(path,*args,**kwargs)
        with patch.object(edits,'atomic_write',side_effect=conflict):
            with self.assertRaises(edits.EditError):edits.move(self.book,self.backups,body)
        self.assertEqual(self.book.read_bytes(),external)

    def test_real_workbook_format_in_memory(self):
        original=server.WORKBOOK.read_bytes();sheets=edits.parse_sheets(original)
        item=edits.master_items(sheets)[0][1];target=next(n for n in edits.registry(sheets['Categories']) if n!=item[3])
        candidate,_=edits.build_move(original,item[0],target,item[3])
        self.assertEqual(server.WORKBOOK.read_bytes(),original)
        with ZipFile(BytesIO(original)) as a,ZipFile(BytesIO(candidate)) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name.startswith(('xl/styles','xl/tables','xl/theme','xl/sharedStrings')):self.assertEqual(a.read(name),b.read(name))

    def test_http_protection_validation_and_status(self):
        old=server.WORKBOOK,server.REVIEW,server.BACKUPS
        server.WORKBOOK,server.REVIEW,server.BACKUPS=self.book,self.root/'review.json',self.backups
        server.REVIEW.write_text(json.dumps({'version':1,'protectedIds':['test.guard']}))
        kept=server.REVIEW.read_bytes()
        http=server.ThreadingHTTPServer(('127.0.0.1',0),partial(server.Handler,directory=str(server.WEB)))
        worker=threading.Thread(target=http.serve_forever,daemon=True);worker.start();base=f'http://127.0.0.1:{http.server_port}'
        def post(body,token=server.TOKEN,origin=base):
            req=Request(base+'/api/items/category',data=json.dumps(body).encode(),headers={'Content-Type':'application/json','Origin':origin,'X-Inventory-Token':token})
            try:
                with urlopen(req) as r:return r.status,json.load(r)
            except HTTPError as e:return e.code,json.load(e)
        try:
            body=self.body();self.assertEqual(post(body,token='bad')[0],403);self.assertEqual(post(body,origin='http://other')[0],403)
            self.assertEqual(post([])[0],400);self.assertEqual(post({**body,'toCategory':'All items'})[0],400)
            self.assertEqual(post(body)[0],200);self.assertEqual(post(body)[0],200)
            self.assertEqual(server.REVIEW.read_bytes(),kept)
            with urlopen(base+'/api/catalog') as r:catalog=json.load(r)
            self.assertIn('Empty category',catalog['categories']);self.assertNotIn('Overlapping family',catalog['categories'])
            with urlopen(base+'/api/category-operations/'+body['operationId']) as r:self.assertEqual(json.load(r)['status'],'committed')
        finally:
            http.shutdown();http.server_close();worker.join();server.WORKBOOK,server.REVIEW,server.BACKUPS=old


if __name__=='__main__':unittest.main()
