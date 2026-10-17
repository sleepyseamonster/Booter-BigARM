"""Protect existing records when additions reach the fixed master formula limit."""
import unittest
import openpyxl
from add_ancient_salvage import reserve_catalog_rows, replace_sheet_row


class RowReservationTests(unittest.TestCase):
    def book(self):
        book = openpyxl.Workbook()
        book.active.title = 'Items'
        book.create_sheet('Craft audit')
        for n in range(6, 11):
            book['Items'].cell(n, 1, f'existing.{n}')
            book['Craft audit'].cell(n, 1, f'existing.{n}')
        return book

    def clear(self, book, n):
        book['Items'].cell(n, 1).value = None
        book['Craft audit'].cell(n, 1).value = None

    def test_tail_is_preferred_with_available_capacity(self):
        book = self.book()
        self.clear(book, 9)
        self.assertEqual(reserve_catalog_rows(book, 10, 2, 12), [11, 12])

    def test_overflow_uses_only_completely_cleared_pairs(self):
        book = self.book()
        for n in [6, 7, 8, 9]: self.clear(book, n)
        book['Craft audit']['J9'] = 'retained audit note'
        book['Items']['AN8'] = '=1'
        book['Items']['B6'] = 'retained name'
        self.assertEqual(reserve_catalog_rows(book, 10, 2, 11), [7, 11])
        self.assertEqual(book['Items']['A10'].value, 'existing.10')

    def test_full_table_recycles_without_extending_capacity(self):
        book = self.book()
        for n in [6, 8, 9]: self.clear(book, n)
        self.assertEqual(reserve_catalog_rows(book, 10, 2, 10), [8, 9])

    def test_insufficient_cleared_pairs_fails_without_mutation(self):
        book = self.book()
        with self.assertRaisesRegex(AssertionError, 'No empty master/audit slots'):
            reserve_catalog_rows(book, 10, 2, 11)
        self.assertEqual(book['Items']['A10'].value, 'existing.10')

    def test_replacing_self_closed_slot_preserves_next_populated_row(self):
        next_row = '<row r="8"><c r="A8" t="inlineStr"><is><t>Keep me</t></is></c></row>'
        raw = '<sheetData><row r="7"/>' + next_row + '</sheetData>'
        replacement = '<row r="7"><c r="A7"/></row>'
        self.assertEqual(replace_sheet_row(raw, 7, replacement, ''), '<sheetData>' + replacement + next_row + '</sheetData>')

    def test_missing_sparse_slot_is_inserted_in_row_order(self):
        raw = '<x:sheetData><x:row r="6"/><x:row r="9"/></x:sheetData>'
        self.assertEqual(replace_sheet_row(raw, 8, '<x:row r="8"/>', 'x:'),
                         '<x:sheetData><x:row r="6"/><x:row r="8"/><x:row r="9"/></x:sheetData>')


if __name__ == '__main__': unittest.main()
